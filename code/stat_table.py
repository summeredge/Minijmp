"""可复用统计表格；str(table) 保留 print_port / TXT 的文本协议。"""
import math
from numbers import Real
from tkinter import ttk, filedialog, messagebox
import tkinter as tk

RESULT_BACKGROUND = '#FFFFFF'
LABEL_CELL_LIMIT = 1000
PNG_MAX_SIDE = 32768
PNG_MAX_PIXELS = 32_000_000  # Pillow RGB 约 4 B/像素；预留编码副本后约 256 MiB。


def configure_result_styles(style):
    """Keep result content white without changing application-wide styles."""
    for name in ('Result.TFrame', 'Result.TLabel', 'Result.TMenubutton'):
        style.configure(name, background=RESULT_BACKGROUND)
        style.map(name, background=[('', RESULT_BACKGROUND)])
    style.configure('Result.Treeview', background=RESULT_BACKGROUND,
                    fieldbackground=RESULT_BACKGROUND, rowheight=26,
                    font=('Microsoft YaHei', 10))


def display_value(value, column):
    if isinstance(value, Real):
        if not math.isfinite(value):
            return '*'
        return str(int(value)) if column in ('N', 'N*', '众数频数') else f'{value:.3f}'
    return '*' if value is None else str(value)


class StatTable(str):
    """标题、列名、二维原值及显示字符串的共同载体。

    兼容 StatTable(text, title, headers, values)。省略 text 时由显示数据生成
    PrettyTable；传入 text（含 PrettyTable 对象）时原样保留其字符串。
    number_formats 为逐列 Excel 格式；不改变原始数值精度或 display。
    """

    def __new__(cls, text=None, title='', headers=(), values=(), *,
                display=None, number_formats=None):
        if isinstance(headers, (str, bytes)):
            raise TypeError('列名必须为字符串序列，不能为单个字符串')
        headers = list(headers)
        if not headers or not all(isinstance(header, str) for header in headers):
            raise ValueError('列名必须为非空的字符串序列')
        if not isinstance(title, str):
            raise TypeError('标题必须为字符串')
        values = list(values)
        if any(isinstance(row, (str, bytes)) for row in values):
            raise ValueError('行数据必须为二维序列')
        values = [list(row) for row in values]
        if any(len(row) != len(headers) for row in values):
            raise ValueError('每行数据的列数必须与列名一致')
        if display is None:
            display = [[display_value(v, h) for h, v in zip(headers, row)]
                       for row in values]
        else:
            display = list(display)
            if any(isinstance(row, (str, bytes)) for row in display):
                raise ValueError('显示数据必须为二维序列')
            display = [list(row) for row in display]
            if len(display) != len(values) or any(len(row) != len(headers) for row in display):
                raise ValueError('显示数据的行列数必须与原值一致')
            if any(not isinstance(cell, str) for row in display for cell in row):
                raise TypeError('显示数据的单元格必须为字符串')
        if number_formats is None:
            number_formats = ['0' if h in ('N', 'N*', '众数频数') else '0.000' for h in headers]
        else:
            if isinstance(number_formats, (str, bytes)):
                raise TypeError('数字格式必须为逐列序列')
            number_formats = list(number_formats)
            if len(number_formats) != len(headers) or any(
                    not isinstance(fmt, str) or not fmt for fmt in number_formats):
                raise ValueError('数字格式必须为与列名等长的非空字符串序列')
        if text is None:
            from prettytable import PrettyTable
            text = PrettyTable(headers)
            text.add_rows(display)
            text.align = 'r'
            text.align[headers[0]] = 'l'
        obj = super().__new__(cls, str(text))
        obj.title = title
        obj.headers = headers
        obj.values = values
        obj.display = display
        obj.number_formats = number_formats
        return obj


def export_excel(table, filename):
    from copy import copy
    from openpyxl import Workbook
    from openpyxl.cell import Cell
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = '统计结果'
    heading_font = Font(name='Microsoft YaHei', bold=True)
    body_font = Font(name='Microsoft YaHei')
    left = Alignment(horizontal='left')
    right = Alignment(horizontal='right')
    styles = []
    for i, fmt in enumerate(table.number_formats):
        heading, text, numeric = (Cell(sheet) for _ in range(3))
        heading.font = heading_font
        text.font = numeric.font = body_font
        heading.alignment = text.alignment = numeric.alignment = right if i else left
        numeric.number_format = fmt
        styles.append((heading, text, numeric))

    def append(values, heading=False):
        cells = []
        for i, value in enumerate(values):
            # 原型只读，复用已绑定的样式；避免每个单元格再查找/注册样式。
            kind = 0 if heading else 2 if isinstance(value, Real) and not isinstance(value, bool) else 1
            cell = copy(styles[i][kind])
            cell.value = value
            if cell.data_type == 'f':  # 标签永远作为文本，避免 Excel 将其解释为公式。
                cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)

    for i, header in enumerate(table.headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = max(12, len(header)*2+2,
            max((len(row[i-1])*2+2 for row in table.display), default=0))
    sheet.freeze_panes = 'B3'
    append([table.title] + [None]*(len(table.headers)-1), heading=True)
    append(table.headers, heading=True)
    for row, shown in zip(table.values, table.display):
        append([v if isinstance(v, Real) and math.isfinite(v) else text
                for v, text in zip(row, shown)])
    book.save(filename)
    book.close()


def export_png(table, filename):
    height = 72 + 36 * (len(table.display) + 1)
    if height > PNG_MAX_SIDE:
        raise ValueError(f'PNG 预计高度 {height:,} 像素，超过安全上限 '
                         f'{PNG_MAX_SIDE:,} 像素。请使用 Excel 导出完整数据。')
    from PIL import Image, ImageDraw, ImageFont
    from matplotlib import font_manager
    path = font_manager.findfont(font_manager.FontProperties(
        family=['Microsoft YaHei', 'SimHei', 'Noto Sans CJK SC', 'DejaVu Sans']))
    face = ImageFont.truetype(path, 20)
    rows = [table.headers] + table.display
    widths = [max(face.getlength(row[i]) for row in rows) + 32
              for i in range(len(table.headers))]
    width = int(max(sum(widths)+32, face.getlength(table.title)+40))
    pixels = width * height
    if width > PNG_MAX_SIDE or pixels > PNG_MAX_PIXELS:
        raise ValueError(f'PNG 预计尺寸 {width:,} × {height:,} 像素，'
                         f'位图约 {pixels*4/1024**2:.1f} MiB，'
                         f'含编码副本约 {pixels*8/1024**2:.1f} MiB，超过安全上限。'
                         '请使用 Excel 导出完整数据。')
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((16, 12), table.title, font=face, fill='#17476b')
    for r, row in enumerate(rows):
        x = 16
        for c, text in enumerate(row):
            offset = widths[c]-16-face.getlength(text) if c and r else 8
            draw.text((x+offset, 60+36*r), text, font=face, fill='#202020')
            x += widths[c]
    draw.line((16, 90, width-16, 90), fill='#9caab5', width=1)
    image.save(filename, format='PNG')
    image.close()


class TableBlock(ttk.Frame):
    def __init__(self, parent, table, *, paginate=False):
        configure_result_styles(ttk.Style(parent))
        super().__init__(parent, padding=8, style='Result.TFrame')
        self.table = table
        self.paged = paginate or (len(table.display)+1)*len(table.headers) > LABEL_CELL_LIMIT
        titlebar = ttk.Frame(self, style='Result.TFrame')
        titlebar.grid(row=0, column=0, sticky='w')
        self.menu_button = button = ttk.Menubutton(titlebar, text='导出', style='Result.TMenubutton')
        button.pack(side='left')
        self.title_label = ttk.Label(titlebar, text=table.title, style='Result.TLabel',
                                    font=('Microsoft YaHei', 11, 'bold'))
        self.title_label.pack(side='left', padx=(8, 0))
        self.menu = tk.Menu(button, tearoff=False)
        for kind in ('Excel', 'PNG'):
            self.menu.add_command(label='导出 '+kind, command=lambda k=kind: self.export(k))
        button.configure(menu=self.menu)
        body = ttk.Frame(self, style='Result.TFrame')
        body.grid(row=1, column=0, sticky='w')
        if self.paged:
            self._setup_pages(body)
            return
        for r, row in enumerate([table.headers] + table.display):
            for c, text in enumerate(row):
                ttk.Label(body, text=text, anchor='e' if c else 'w', style='Result.TLabel',
                          font=('Microsoft YaHei', 10, 'bold' if r == 0 else 'normal')).grid(
                    row=r, column=c, sticky='ew', padx=10, pady=3)

    def _setup_pages(self, body):
        # 同时限制行数与单页单元格数，不向 Tcl 装载整张表。
        self.page_size = min(100, max(1, LABEL_CELL_LIMIT // len(self.table.headers)))
        self.page_count = max(1, math.ceil(len(self.table.display) / self.page_size))
        self.page = 1
        self.page_number = tk.StringVar(self, '1')
        columns = [str(c) for c in range(len(self.table.headers))]
        body.configure(width=640, height=400)
        body.grid_propagate(False)
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(body, columns=columns, show='headings',
                                 style='Result.Treeview')
        for c, header in enumerate(self.table.headers):
            self.tree.heading(str(c), text=header)
            self.tree.column(str(c), width=max(100, min(260, len(header)*14+20)),
                             minwidth=40, stretch=False, anchor='e' if c else 'w')
        vertical = ttk.Scrollbar(body, orient='vertical', command=self.tree.yview)
        horizontal = ttk.Scrollbar(body, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        self.tree.grid(row=0, column=0, sticky='nsew')
        vertical.grid(row=0, column=1, sticky='ns')
        horizontal.grid(row=1, column=0, sticky='ew')
        navigation = ttk.Frame(self, style='Result.TFrame')
        navigation.grid(row=2, column=0, sticky='w', pady=(4, 0))
        for text, destination in [('首页', lambda: 1), ('上页', lambda: self.page-1)]:
            ttk.Button(navigation, text=text, width=5,
                       command=lambda target=destination: self.show_page(target())).pack(side='left')
        ttk.Label(navigation, text='第', style='Result.TLabel').pack(side='left')
        self.page_input = ttk.Spinbox(navigation, from_=1, to=self.page_count, width=6,
                                     textvariable=self.page_number,
                                     command=lambda: self.show_page(self.page_number.get()))
        self.page_input.pack(side='left')
        self.page_input.bind('<Return>', lambda event: self.show_page(self.page_number.get()))
        ttk.Label(navigation, text=f'/ {self.page_count} 页', style='Result.TLabel').pack(side='left')
        for text, destination in [('下页', lambda: self.page+1), ('末页', lambda: self.page_count)]:
            ttk.Button(navigation, text=text, width=5,
                       command=lambda target=destination: self.show_page(target())).pack(side='left')
        self.page_status = ttk.Label(navigation, style='Result.TLabel')
        self.page_status.pack(side='left', padx=8)
        for sequence, destination in [('<Prior>', lambda: self.page-1),
                                      ('<Next>', lambda: self.page+1),
                                      ('<Control-Home>', lambda: 1),
                                      ('<Control-End>', lambda: self.page_count)]:
            self.tree.bind(sequence, lambda event, target=destination: self.show_page(target()))
        self.show_page(1)

    def show_page(self, page):
        try:
            page = int(page)
        except ValueError:
            self.page_number.set(str(self.page))
            return 'break'
        self.page = max(1, min(self.page_count, page))
        self.page_number.set(str(self.page))
        children = self.tree.get_children()
        if children:
            self.tree.delete(*children)
        start = (self.page-1)*self.page_size
        stop = min(start+self.page_size, len(self.table.display))
        for index in range(start, stop):
            self.tree.insert('', 'end', iid=str(index), values=self.table.display[index])
        self.tree.yview_moveto(0)
        self.page_status.configure(text=f'{start+1 if stop else 0}–{stop} / '
                                       f'{len(self.table.display)} 行')
        return 'break'

    def destroy(self):
        super().destroy()
        self.table = None

    def export(self, kind):
        extension = '.xlsx' if kind == 'Excel' else '.png'
        filename = filedialog.asksaveasfilename(parent=self, defaultextension=extension,
            filetypes=[(kind, '*'+extension)])
        if not filename:
            return
        try:
            (export_excel if kind == 'Excel' else export_png)(self.table, filename)
        except Exception as error:
            messagebox.showerror('导出失败', str(error), parent=self)

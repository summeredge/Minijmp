from describe import describe
from stat_table import StatTable

def test_structured_fallback():
    lines = []
    describe([-2, 1, None, 4], print_out=True, print_port=lines.append)
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert len(tables) == 3
    assert tables[0].values[0][1:3] == [3, 1]
    assert tables[0].display[0][3] == '1.000'
    assert '|' in str(tables[0])
import ast
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import tkinter as tk
import pytest
from openpyxl import load_workbook
from PIL import Image
import stat_table
import describe as description
import basic_dialogs
import MiniJMP
from test_g1_text_output import viewer


def reports():
    lines = []
    description.describe([-5, None, -2, 4, 7], print_out=True, print_port=lines.append)
    return lines


def test_real_exports(tmp_path):
    for i, table in enumerate(x for x in reports() if isinstance(x, StatTable)):
        excel = tmp_path / f'中文{i}.xlsx'
        png = tmp_path / f'中文{i}.png'
        stat_table.export_excel(table, excel)
        stat_table.export_png(table, png)
        book = load_workbook(excel)
        sheet = book.active
        assert sheet.cell(1,1).value == table.title
        assert [c.value for c in sheet[2]] == table.headers
        for r, row in enumerate(table.values, 3):
            for c, value in enumerate(row, 1):
                actual = sheet.cell(r,c)
                if isinstance(value, (int,float)) and value == value:
                    assert actual.value == pytest.approx(value)
                    assert actual.data_type == 'n'
                else:
                    assert actual.value == table.display[r-3][c-1]
        book.close()
        with Image.open(png) as image:
            assert image.format == 'PNG'
            assert image.height == 72+36*(len(table.values)+1)
            assert image.getbbox()


def test_text_contract_order_copy_save_clear(viewer, tmp_path):
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    lines = reports()
    for line in lines:
        MiniJMP.Minijmp.print(app, line)
    MiniJMP.Minijmp.print(app, '普通报告\n+---+\n| a |')
    expected = ''.join(str(x)+'\n' for x in lines)+'普通报告\n+---+\n| a |\n'
    viewer.main.update()
    assert viewer.T.get('1.0','end-1c') == expected
    assert len(viewer.table_blocks) == 3
    assert all(b.winfo_ismapped() for b in viewer.table_blocks)
    viewer.select_all()
    assert viewer.T.get('sel.first','sel.last') == expected
    path=tmp_path/'报告.txt'
    viewer.save(path)
    assert path.read_text(encoding='utf-8') == expected
    assert 'table_ascii' not in viewer.T.tag_names(viewer.T.search('普通报告','1.0'))
    viewer.clear()
    assert viewer.table_blocks == []
    assert viewer.T.get('1.0','end-1c') == ''


@pytest.mark.parametrize('kind', ['Excel', 'PNG'])
def test_cancel_error_menu(viewer, monkeypatch, tmp_path, kind):
    table = next(x for x in reports() if isinstance(x, StatTable))
    viewer.add_table(table)
    block = viewer.table_blocks[0]
    assert [block.menu.entrycget(i,'label') for i in range(2)] == ['导出 Excel','导出 PNG']
    save=Mock(return_value='')
    monkeypatch.setattr(stat_table.filedialog,'asksaveasfilename',save)
    exporter=Mock(side_effect=OSError('拒绝写入'))
    monkeypatch.setattr(stat_table,'export_excel' if kind == 'Excel' else 'export_png',exporter)
    error=Mock()
    monkeypatch.setattr(stat_table.messagebox,'showerror',error)
    block.export(kind)
    exporter.assert_not_called()
    save.return_value=str(tmp_path/'结果.xlsx')
    block.menu.invoke(0 if kind == 'Excel' else 1)
    error.assert_called_once()
    assert str(table) in viewer.T.get('1.0','end-1c')


def test_missing_wide_long_negative_and_formula(tmp_path, viewer):
    table=StatTable('旧文本', '中文完整导出', ['变量']+[f'统计{i}' for i in range(30)],
        [['=危险标签']+[-1.234567, None, float('nan')]+list(range(27)) for _ in range(100)])
    viewer.add_table(table)
    viewer.main.update()
    block = viewer.table_blocks[0]
    assert block.paged and block.page_count > 1
    assert block.tree.xview()[1] < 1
    assert block.tree.yview()[1] < 1
    stat_table.export_png(table,tmp_path/'宽长.png')
    with Image.open(tmp_path/'宽长.png') as image:
        assert image.width > 2500 and image.height > 3600
    stat_table.export_excel(table,tmp_path/'宽长.xlsx')
    book=load_workbook(tmp_path/'宽长.xlsx')
    assert book.active.max_row == 102 and book.active.max_column == 31
    assert book.active['A3'].data_type == 's'
    assert book.active['C3'].value == '*'
    book.close()


@pytest.mark.parametrize('multiple',[False, True])
def test_dialog_structured_route(multiple):
    import pandas as pd
    var=lambda value: SimpleNamespace(get=lambda:value)
    lines=[]
    dialog=SimpleNamespace(df=pd.DataFrame({'温度':[-2,None,4], '压力':[1,3,8]}),
        xvar=var('温度'),pct=var('100,50,0'),grpcols=['温度','压力'],
        names=list(description.STAT_LABELS),picks={key:var(True) for key in description.STAT_LABELS},
        app=SimpleNamespace(print=lines.append),error=Mock())
    (basic_dialogs.MultiDescribeDialog if multiple else basic_dialogs.DescribeDialog).apply(dialog)
    tables=[x for x in lines if isinstance(x,StatTable)]
    assert len(tables) == (2 if multiple else 3)
    assert len(tables[0].values) == (2 if multiple else 1)
    assert tables[0].values[0][1:3] == [2,1]
    dialog.error.assert_not_called()


def test_source_scope_and_no_ascii_parser():
    root=Path(__file__).resolve().parents[1]
    tree=ast.parse((root/'code/describe.py').read_text(encoding='utf-8'))
    assert sum(isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='StatTable'
               for n in ast.walk(tree)) == 2
    for filename in ('basic_dialogs.py','describe.py','txt_output.py','MiniJMP.py'):
        source=(root/'code'/filename).read_text(encoding='utf-8')
        assert 're.findall' not in source
    assert 'Text.get' not in (root/'code/stat_table.py').read_text(encoding='utf-8')

def test_native_copy_includes_hidden_ascii(viewer):
    table = next(x for x in reports() if isinstance(x, StatTable))
    viewer.add_table(table)
    viewer.add_txt('结束\n')
    viewer.select_all()
    copied=[]
    interp=viewer.T.tk
    interp.call('rename','clipboard','g7_original_clipboard')
    interp.createcommand('clipboard',lambda *args: copied.append(args) or '')
    try:
        interp.call('tk_textCopy', str(viewer.T))
    finally:
        interp.deletecommand('clipboard')
        interp.call('rename','g7_original_clipboard','clipboard')
    assert copied[-1][-1] == str(table)+'\n结束\n'


def test_reopen_and_repeated_output(viewer, monkeypatch):
    factory=Mock(return_value=viewer)
    monkeypatch.setattr(MiniJMP,'txt_viewer',factory)
    app=SimpleNamespace(table=SimpleNamespace(tOut=None))
    table=next(x for x in reports() if isinstance(x,StatTable))
    MiniJMP.Minijmp.print(app,table,end='\n')
    app.table.tOut=viewer
    MiniJMP.Minijmp.print(app,table,end='\n')
    assert viewer.T.get('1.0','end-1c') == (str(table)+'\n')*2
    factory.assert_called_once()


def test_display_missing_precision():
    assert stat_table.display_value(None,'均值') == '*'
    assert stat_table.display_value(float('inf'),'均值') == '*'
    assert stat_table.display_value(-1.23456,'均值') == '-1.235'
    assert stat_table.display_value(2,'N*') == '2'

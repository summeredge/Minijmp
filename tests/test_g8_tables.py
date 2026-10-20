"""G8：公共接口、真实 Tk 布局和 G7 文本/导出契约。"""
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook
from PIL import Image
from prettytable import PrettyTable

import basic_dialogs
import describe as description
import stat_table
from MiniJMP import Minijmp
from stat_table import StatTable, export_excel, export_png
from test_g1_text_output import viewer


CASES = ['function', 'multiple', 'supplemental', 'percentiles', 'single_dialog', 'multi_dialog']


def report(case):
    lines = []
    data = [-5, None, -2, 4, 7]
    result = description.describe(data)
    if case == 'function':
        description.describe(data, print_out=True, print_port=lines.append)
    elif case == 'multiple':
        description.print_descriptive({'温度': result, '压力': description.describe([1, 3, 8])},
                                      print_port=lines.append, missing={'温度': 1, '压力': 0})
    elif case == 'supplemental':
        description.print_descriptive({'温度': result}, print_port=lines.append,
                                      selected=['Variance', 'Sum', 'N of Mode'])
    elif case == 'percentiles':
        description.print_percentiles([100, 50, 0], [7, 1, -5], print_port=lines.append)
    else:
        var = lambda value: SimpleNamespace(get=lambda: value)
        dialog = SimpleNamespace(df=pd.DataFrame({'温度': data, '压力': [1, 3, 8, 9, 10]}),
            xvar=var('温度'), pct=var('100,50,0'), grpcols=['温度', '压力'],
            names=list(description.STAT_LABELS),
            picks={key: var(True) for key in description.STAT_LABELS},
            app=SimpleNamespace(print=lines.append), error=Mock())
        cls = basic_dialogs.DescribeDialog if case == 'single_dialog' else basic_dialogs.MultiDescribeDialog
        cls.apply(dialog)
        dialog.error.assert_not_called()
    return lines


@pytest.mark.parametrize('case', CASES)
def test_g7_text_baseline_and_visible_headings(viewer, tmp_path, case):
    baseline = json.loads((Path(__file__).parent / 'fixtures/g8_g7_text.json').read_text(encoding='utf-8'))
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    lines = report(case)
    for line in lines:
        Minijmp.print(app, line)
    viewer.main.update()
    assert viewer.T.get('1.0', 'end-1c') == baseline[case]
    target = tmp_path / '原报告.txt'
    viewer.save(target)
    assert target.read_text(encoding='utf-8') == baseline[case]
    visible = ''.join(item[1] for item in viewer.T.dump('1.0', 'end-1c', text=True)
                      if not any(viewer.T.tag_cget(tag, 'elide') == '1'
                                 for tag in viewer.T.tag_names(item[2])))
    for block in viewer.table_blocks:
        assert block.table.title not in visible
        assert block.title_label.cget('text') == block.table.title


def test_public_interface_and_legacy_text(capsys):
    legacy = PrettyTable(['变量', '比例'])
    legacy.add_row(['温度', '12.3%'])
    table = StatTable(legacy, '中文标题', ['变量', '比例'], [['温度', .123]],
                      display=[['温度', '12.3%']], number_formats=['General', '0.0%'])
    assert isinstance(table, str)
    assert str(table) == str(legacy)
    assert table.values == [['温度', .123]]
    assert table.display == [['温度', '12.3%']]
    print(table)
    assert capsys.readouterr().out == str(legacy) + '\n'
    automatic = StatTable(title='自动表', headers=['变量', '均值'], values=[['压力', -1.23456]])
    assert '-1.235' in str(automatic) and '压力' in str(automatic)


@pytest.mark.parametrize('kwargs', [
    {'headers': []},
    {'headers': '甲乙'},
    {'title': None},
    {'values': [[1]]},
    {'values': [[1, 2, 3]]},
    {'values': ['ab']},
    {'display': []},
    {'display': [['缺一列']]},
    {'display': ['ab']},
    {'display': [[1, 2]]},
    {'display': [['a', 'b'], ['c', 'd']]},
    {'number_formats': ['0']},
    {'number_formats': '00'},
    {'number_formats': ['', '0']},
])
def test_public_interface_rejects_truncated_data(kwargs):
    arguments = dict(title='错误数据', headers=['甲', '乙'], values=[[1, 2]])
    arguments.update(kwargs)
    with pytest.raises((ValueError, TypeError)):
        StatTable(**arguments)


@pytest.mark.parametrize('rows', [[], [['中文', None, float('nan'), float('inf'), -1.23456789]]])
def test_empty_missing_and_real_exports(tmp_path, viewer, rows):
    table = StatTable(title='空表或缺失值', headers=['变量', '空', '缺失', '无穷', '估计值'], values=rows)
    viewer.add_table(table)
    viewer.main.update()
    assert viewer.table_blocks[0].winfo_ismapped()
    export_excel(table, tmp_path / '结果.xlsx')
    book = load_workbook(tmp_path / '结果.xlsx')
    assert book.active.max_row == 2 + len(rows)
    assert book.active.max_column == 5
    if rows:
        assert [book.active.cell(3, col).value for col in range(2, 5)] == ['*'] * 3
        assert book.active['E3'].value == -1.23456789
        assert book.active['E3'].number_format == '0.000'
    book.close()
    export_png(table, tmp_path / '结果.png')
    with Image.open(tmp_path / '结果.png') as image:
        assert image.height == 72 + 36 * (len(rows) + 1)
        assert image.crop((0, 0, image.width, 55)).getextrema() != ((255, 255),) * 3


def test_custom_excel_format_and_complete_png(tmp_path):
    table = StatTable(title='完整导出', headers=['变量', '比例', 'N'],
        values=[['中文', .12345678, 12]], display=[['中文', '12.35%', '12']],
        number_formats=['General', '0.00%', '0'])
    export_excel(table, tmp_path / '格式.xlsx')
    book = load_workbook(tmp_path / '格式.xlsx')
    assert book.active['B3'].value == .12345678
    assert book.active['B3'].data_type == 'n'
    assert book.active['B3'].number_format == '0.00%'
    assert book.active['C3'].number_format == '0'
    book.close()
    # 最后一个单元格也必须参与绘制，不能受 GUI 视口裁切。
    table = StatTable(title='完整导出', headers=['变量', '末列'],
                      values=[['', ''] for _ in range(99)] + [['', '最后单元格']])
    export_png(table, tmp_path / '长表.png')
    with Image.open(tmp_path / '长表.png') as image:
        assert image.height == 3708
        assert image.crop((0, 60 + 36 * 100, image.width, 60 + 36 * 101)).getextrema() != ((255, 255),) * 3


@pytest.mark.parametrize('geometry', ['480x320', '1100x800'])
def test_menu_position_independent_of_columns_and_scroll(viewer, geometry):
    viewer.main.geometry(geometry)
    for columns in (2, 10, 31):
        viewer.add_table(StatTable(title='同级表格', headers=[f'列{i}' for i in range(columns)],
                                   values=[list(range(columns)) for _ in range(8)]))
    viewer.main.update()
    assert len({block.menu_button.winfo_rootx() for block in viewer.table_blocks}) == 1
    for block in viewer.table_blocks:
        assert block.menu_button.winfo_x() < block.title_label.winfo_x()
        assert block.menu_button.winfo_width() > 20
    assert viewer.T.yview()[1] < 1
    viewer.T.yview_moveto(1)
    viewer.main.update()
    assert viewer.T.xview()[1] < 1
    viewer.T.xview_moveto(1)
    viewer.main.update()
    assert viewer.T.xview()[0] > 0 and viewer.T.yview()[0] > 0


def test_only_adjacent_same_heading_hidden_and_clear_resets(viewer):
    table = StatTable('原表', '标题', ['列'], [[1]])
    viewer.add_txt('标题\n这里是说明，必须保留\n')
    viewer.add_table(table)
    assert not viewer.T.tag_ranges('table_heading')
    viewer.add_txt('其他章节\n')
    viewer.add_table(table)
    assert not viewer.T.tag_ranges('table_heading')
    viewer.add_txt('标题\n\n')
    viewer.add_table(table)
    ranges = viewer.T.tag_ranges('table_heading')
    assert viewer.T.get(*ranges) == '标题\n'
    viewer.add_txt('普通输出\n')
    index = viewer.T.search('普通输出', '1.0')
    assert 'table_heading' not in viewer.T.tag_names(index)
    viewer.clear()
    viewer.add_table(table)
    assert not viewer.T.tag_ranges('table_heading')
    assert viewer.T.get('1.0', 'end-1c') == '原表\n'


def test_consecutive_tables_do_not_hide_previous_window(viewer):
    table = StatTable('标题', '标题', ['列'], [[1]])
    for _ in range(3):
        viewer.add_table(table)
    viewer.main.update()
    assert len(viewer.table_blocks) == 3
    assert all(block.winfo_ismapped() for block in viewer.table_blocks)
    assert not viewer.T.tag_ranges('table_heading')


def test_copy_hidden_heading_preserves_original_report(viewer):
    viewer.add_txt('标题\n')
    viewer.add_table(StatTable('原文本', '标题', ['列'], [[1]]))
    viewer.select_all()
    copied = []
    interp = viewer.T.tk
    interp.call('rename', 'clipboard', 'g8_original_clipboard')
    interp.createcommand('clipboard', lambda *args: copied.append(args) or '')
    try:
        interp.call('tk_textCopy', str(viewer.T))
    finally:
        interp.deletecommand('clipboard')
        interp.call('rename', 'g8_original_clipboard', 'clipboard')
    assert copied[-1][-1] == '标题\n原文本\n'


@pytest.mark.parametrize('kind,extension', [('Excel', '.xlsx'), ('PNG', '.png')])
def test_menu_writes_real_file(viewer, monkeypatch, tmp_path, kind, extension):
    table = StatTable(title='菜单导出', headers=['中文', 'N'], values=[['末行', 23]])
    viewer.add_table(table)
    filename = tmp_path / ('菜单' + extension)
    monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kwargs: str(filename))
    error = Mock()
    monkeypatch.setattr(stat_table.messagebox, 'showerror', error)
    viewer.table_blocks[0].menu.invoke(0 if kind == 'Excel' else 1)
    error.assert_not_called()
    if kind == 'Excel':
        book = load_workbook(filename)
        assert book.active['A3'].value == '末行' and book.active['B3'].value == 23
        book.close()
    else:
        with Image.open(filename) as image:
            assert image.format == 'PNG' and image.height == 144


def test_titlebar_at_double_dpi(viewer):
    interp = viewer.T.tk
    original = interp.call('tk', 'scaling')
    try:
        interp.call('tk', 'scaling', 2.66667)
        for columns in (2, 31):
            viewer.add_table(StatTable(title='高 DPI 中文标题', headers=[f'列{i}' for i in range(columns)],
                                       values=[list(range(columns))]))
        viewer.main.update()
        assert len({block.menu_button.winfo_rootx() for block in viewer.table_blocks}) == 1
        for block in viewer.table_blocks:
            assert block.menu_button.winfo_x() + block.menu_button.winfo_width() <= block.title_label.winfo_x()
            assert block.title_label.winfo_width() >= block.title_label.winfo_reqwidth()
    finally:
        interp.call('tk', 'scaling', original)

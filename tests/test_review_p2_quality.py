"""P2 质量工具反例：整数比例和人工分组独立参考。"""
from fractions import Fraction
from io import StringIO

import numpy as np
import pandas as pd
import pytest

from quality_tools import multi_vari, pareto_chart
from quality_tools_plots import plot_quality


@pytest.mark.parametrize('grouped', [False, True])
@pytest.mark.parametrize('threshold,expected', [(55, [55, 30, 15]),
                                             (np.nextafter(55., 0.), [55, 45]),
                                             (60, [55, 30, 15])])
def test_pareto_strict_threshold(threshold, expected, grouped):
    frame = pd.DataFrame({'c': ['A', 'B', 'C'], 'n': [55, 30, 15], 'g': ['G']*3})
    table = next(iter(pareto_chart(frame, 'c', 'n', group='g' if grouped else None,
                                  threshold=threshold)['tables'].values()))
    assert table['频数'].tolist() == expected
    assert table['类别'].tolist() == (['A', 'B', '其他（合并）'] if len(expected) == 3 else ['A', '其他（合并）'])
    np.testing.assert_allclose(table['百分比'], [float(Fraction(v, 100)*100) for v in expected])
    assert table['累计百分比'].iloc[-1] == 100


@pytest.mark.parametrize('grouped', [False, True])
@pytest.mark.parametrize('merge', [False, True])
def test_pareto_big_integer_counts_and_export(grouped, merge):
    frame = pd.DataFrame({'c': ['A', 'A', 'B', 'C'], 'n': [2**62, 2**62, 2, 1], 'g': ['G']*4})
    table = next(iter(pareto_chart(frame, 'c', 'n', group='g' if grouped else None,
                                  threshold=None, max_categories=2 if merge else None)['tables'].values()))
    expected = [2**63, 3] if merge else [2**63, 2, 1]
    assert table['频数'].tolist() == expected
    assert table['类别'].tolist() == (['A', '其他（合并）'] if merge else ['A', 'B', 'C'])
    assert sum(table['频数']) == 2**63+3
    np.testing.assert_allclose(table['百分比'], [float(Fraction(v, 2**63+3)*100) for v in expected])
    assert table['累计百分比'].iloc[-1] == 100
    exported = pd.read_csv(StringIO(table.to_csv(index=False)), dtype={'频数': str})
    assert exported['频数'].tolist() == [str(v) for v in expected]


def test_pareto_arbitrary_integer_and_unrepresentable_percentage():
    floating = pd.DataFrame({'c': ['A'], 'n': [float(2**63)]})
    assert pareto_chart(floating, 'c', 'n')['tables']['分组 全部']['频数'].tolist() == [2**63]
    frame = pd.DataFrame({'c': ['A', 'A', 'B'], 'n': pd.Series([10**400, str(10**400), 10**400], dtype=object)})
    table = pareto_chart(frame, 'c', 'n', threshold=None)['tables']['分组 全部']
    assert table['频数'].tolist() == [2*10**400, 10**400]
    np.testing.assert_allclose(table['百分比'], [200/3, 100/3])
    frame.loc[2, 'n'] = 1
    with pytest.raises(ValueError, match='百分比可表达范围'):
        pareto_chart(frame, 'c', 'n', threshold=None)


@pytest.mark.parametrize('name', ['N', 'SD', '均值', '最小值', '最大值', '极差'])
def test_multi_vari_reserved_factor_names(name, tmp_path):
    frame = pd.DataFrame({name: ['B', 'B', 'A', 'A'], f'{name}（因子）': [0, 1, 0, 1], 'y': [1., 3., 5., 7.]})
    original = frame.copy(deep=True)
    result = multi_vari(frame, 'y', [name, f'{name}（因子）'])
    first, second = result['factors']
    assert first == f'{name}（因子）*'
    assert second == f'{name}（因子）'
    table = result['tables']['组合统计']
    assert table[first].tolist() == ['B', 'B', 'A', 'A']
    assert table['N'].tolist() == [1]*4
    assert table['均值'].tolist() == [1, 3, 5, 7]
    assert table['SD'].isna().all()
    assert table['极差'].tolist() == [0]*4
    assert result['coverage'] == 1
    pd.testing.assert_series_equal(result['observations'][first], frame[name], check_names=False)
    pd.testing.assert_frame_equal(frame, original)
    figure = plot_quality(result)
    np.testing.assert_array_equal(figure.axes[0].lines[0].get_ydata(), [1, 5])
    np.testing.assert_array_equal(figure.axes[0].lines[1].get_ydata(), [3, 7])
    figure.savefig(tmp_path/f'multi-vari-{name}.png')


def test_pareto_large_integer_keeps_exact_digits_in_both_excel_exports(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import Mock
    from openpyxl import load_workbook
    from quality_tools_dialog import QualityToolsDialog
    from stat_table import export_excel
    frame = pd.DataFrame({'c': ['A', 'A', 'B'], 'n': [2**62, 2**62+1, 1]})
    table = pareto_chart(frame, 'c', 'n', threshold=None)['tables']['分组 全部']
    expected = 2**63+1
    direct = tmp_path / 'direct.xlsx'
    monkeypatch.setattr('quality_tools_dialog.filedialog.asksaveasfilename', lambda **kw: str(direct))
    dialog = SimpleNamespace(main=None, current_table=lambda: ('频数', table), error=Mock())
    QualityToolsDialog.export(dialog)
    dialog.error.assert_not_called()
    book = load_workbook(direct)
    assert book.active['B2'].value == str(expected)
    assert book.active['B2'].data_type == 's'
    book.close()
    report = tmp_path / 'report.xlsx'
    export_excel(QualityToolsDialog.stat_table('频数', table), report)
    book = load_workbook(report)
    assert book.active['B3'].value == str(expected)
    assert book.active['B3'].data_type == 's'
    book.close()
    assert table['频数'].tolist() == [expected, 1]
    factor_result = multi_vari(pd.DataFrame({'频数': ['A', 'B'], 'y': [1, 2]}), 'y', ['频数'])
    shown = QualityToolsDialog.stat_table('组合统计', factor_result['tables']['组合统计'])
    assert [row[0] for row in shown.values] == ['A', 'B']

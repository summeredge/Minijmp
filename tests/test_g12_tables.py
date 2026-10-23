"""G12：遗漏的区间报告及全部迁移模块在混合报告中的契约。"""
import ast
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from PIL import Image

import basic_dialogs
import stat_table
import utilities
from MiniJMP import Minijmp
from stat_table import StatTable
from test_g1_text_output import viewer
from test_g3_computation_contracts import computation_dump
from source_baselines import read_source_baseline

ROOT = Path(__file__).resolve().parents[1]
BASE = 'c0c2c2aa316e7e0f236e551bf4cd2f60f2e30df3'
DATA = [1.123456789, 2., 3., 5., 7., 9.]


@pytest.fixture(scope='module')
def before():
    source = read_source_baseline(BASE, 'code/utilities.py')
    module = ModuleType('g12_before_utilities')
    exec(compile(source, '<G12-before:utilities>', 'exec'), module.__dict__)
    return module


@pytest.mark.parametrize('alpha', [.05, .1])
@pytest.mark.parametrize('scale', [1., 1e-80, 1e80])
def test_interval_report_keeps_txt_computation_and_exports(before, alpha, scale, tmp_path):
    data = [value * scale for value in DATA]
    old, lines = [], []
    expected = before.mean_std_CIs(data, alpha=alpha, print_out=True, print_port=old.append)
    result = utilities.mean_std_CIs(data, alpha=alpha, print_out=True, print_port=lines.append)
    assert result == expected
    assert '\n'.join(lines) == '\n'.join(old)
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert len(tables) == 1
    table = tables[0]
    assert table.title == '置信区间'
    for row, key, point in zip(table.values, ['mean', 'std'],
                               [np.mean(data), np.std(data, ddof=1)]):
        assert row[1] == pytest.approx(point, rel=1e-14, abs=0)
        assert row[2:] == [*result[key+' 2side'], *result[key+' 1side']]
    silent = []
    assert utilities.mean_std_CIs(data, alpha=alpha, print_port=silent.append) == expected
    assert silent == []
    excel, png = tmp_path/'区间.xlsx', tmp_path/'区间.png'
    stat_table.export_excel(table, excel)
    stat_table.export_png(table, png)
    book = load_workbook(excel)
    assert book.active['A1'].value == table.title
    assert [cell.value for cell in book.active[2]] == table.headers
    assert book.active.max_row == 4 and book.active.max_column == 6
    for r, row in enumerate(table.values, 3):
        for c, raw in enumerate(row[1:], 2):
            cell = book.active.cell(r, c)
            assert cell.data_type == 'n'
            assert cell.value == pytest.approx(raw, rel=1e-14, abs=0)
            assert cell.number_format == table.number_formats[c-1]
    book.close()
    with Image.open(png) as image:
        assert image.height == 180
        assert image.crop((0, 132, image.width, 162)).getextrema() != ((255, 255),)*3


def test_utilities_change_is_only_interval_report_output():
    old = read_source_baseline(BASE, 'code/utilities.py')
    current = (ROOT/'code/utilities.py').read_text(encoding='utf-8')
    assert computation_dump(current, 'mean_std_CIs') == computation_dump(old, 'mean_std_CIs')
    trees = [ast.parse(source) for source in (old, current)]
    for tree in trees:
        # G17 仅放行已另有数据一致性及热点回归约束的数字清洗函数。
        tree.body = [n for n in tree.body if not (
            isinstance(n, ast.FunctionDef) and n.name == 'number_2Dlist')]
        function = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                        and n.name == 'mean_std_CIs')
        function.body = [n for n in function.body if not (
            isinstance(n, ast.If) and isinstance(n.test, ast.Name) and n.test.id == 'print_out')]
    assert ast.dump(trees[0]) == ast.dump(trees[1])


def test_ci_dialog_routes_structured_table_and_keeps_original_text(before, viewer):
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    dialog = SimpleNamespace(remember_columns=Mock(), df=pd.DataFrame({'测量': DATA + [None, '无效']}),
        xvar=SimpleNamespace(get=lambda: '测量'), alpha=SimpleNamespace(get=lambda: .1),
        app=SimpleNamespace(print=lambda line: Minijmp.print(app, line)))
    basic_dialogs.CIDialog.apply(dialog)
    old = []
    data = before.number_list(dialog.df['测量'], col_name='测量', print_out=True, print_port=old.append)
    before.mean_std_CIs(data, alpha=.1, print_out=True, print_port=old.append)
    assert viewer.T.get('1.0', 'end-1c') == ''.join(str(line)+'\n' for line in old)
    assert len(viewer.table_blocks) == 1


def mixed_reports():
    from test_g7_tables import reports as describe_reports
    from test_g9_tables import report as hypothesis_report
    from test_g10_tables import reports as model_reports
    from test_g11_quality_tables import report as quality_report
    lines = ['混合统计报告：中文 μ、ρ、α 与 =标签'] + describe_reports()
    utilities.mean_std_CIs(DATA, print_out=True, print_port=lines.append)
    for case in ['t', 'prop1', 'sd', 'contingency', 'gof_summary', 'one', 'two_inter', 'ti']:
        lines.extend(hypothesis_report(case))
    models = model_reports()
    for case in ['ols', 'ortho', 'normality', 'pearson', 'resid', 'compare', 'spearman',
                 'kendall', 'hoeffding', 'spearman_matrix', 'multi_fit', 'multi_fit_orig']:
        models[case](lines.append)
    for case in [('direct', 'emp', True, True), ('nested', None),
                 ('cpk', True, 'unbalanced', True), ('spc', 'NP')]:
        lines.extend(quality_report(case)[1])
    # 保留有边框的普通字符串及错误消息，不推断为可导出统计表。
    lines.extend(['\n+---+\n| 纯文本 |\n+---+', 'ERROR：输入无效', '报告结束'])
    return lines


def test_all_modules_mixed_report_copy_save_scroll_clear_close_reopen(viewer, tmp_path, monkeypatch):
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    lines = mixed_reports()
    tables = [line for line in lines if isinstance(line, StatTable)]
    for line in lines:
        Minijmp.print(app, line)
    viewer.main.geometry('480x320')
    viewer.main.update()
    expected = ''.join(str(line)+'\n' for line in lines)
    assert viewer.T.get('1.0', 'end-1c') == expected
    assert [id(block.table) for block in viewer.table_blocks] == list(map(id, tables))
    assert len(tables) > 50
    window_indices = {entry[1]: entry[2] for entry in viewer.T.dump('1.0', 'end', window=True)}
    menu_positions = []
    export_error = Mock()
    monkeypatch.setattr(stat_table.messagebox, 'showerror', export_error)
    for number, block in enumerate(viewer.table_blocks):
        # Text 仅为可见内嵌窗口更新屏幕坐标，先滚到该表再比较菜单位置。
        viewer.T.see(window_indices[str(block)])
        viewer.T.xview_moveto(0)
        viewer.main.update()
        assert block.winfo_ismapped()
        menu_positions.append(block.menu_button.winfo_rootx())
        assert block.menu_button.winfo_x() > block.title_label.winfo_x()
        assert block.title_label.cget('text') == block.table.title
        assert block.menu.index('end') == 2
        body = block.winfo_children()[1]
        for label in body.winfo_children():
            column = label.grid_info()['column']
            assert str(label.cget('anchor')) == ('e' if column else 'w')
            assert 'Microsoft YaHei' in str(label.cget('font'))
        for menu_index, suffix in [(0, '.xlsx'), (1, '.png')]:
            target = tmp_path / f'{number}{suffix}'
            monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename',
                                lambda **kwargs: str(target))
            block.menu.invoke(menu_index)
            export_error.assert_not_called()
            if menu_index == 0:
                book = load_workbook(target)
                assert book.active['A1'].value == block.table.title
                assert book.active.max_row == len(block.table.values)+2
                assert book.active.max_column == len(block.table.headers)
                book.close()
            else:
                with Image.open(target) as image:
                    assert image.height == 72+36*(len(block.table.values)+1)
                    image.verify()
    assert len(set(menu_positions)) == 1
    viewer.T.yview_moveto(0)
    viewer.main.update()
    assert viewer.T.yview()[1] < 1
    widest = max(viewer.table_blocks, key=lambda b: b.winfo_reqwidth())
    index = next(entry[2] for entry in viewer.T.dump('1.0', 'end', window=True)
                 if entry[1] == str(widest))
    viewer.T.see(index)
    viewer.main.update()
    assert viewer.T.xview()[1] < 1
    viewer.T.xview_moveto(1)
    viewer.main.update()
    assert viewer.T.xview()[0] > 0
    viewer.select_all()
    copied = []
    interp = viewer.T.tk
    interp.call('rename', 'clipboard', 'g12_original_clipboard')
    interp.createcommand('clipboard', lambda *args: copied.append(args) or '')
    try:
        interp.call('tk_textCopy', str(viewer.T))
    finally:
        interp.deletecommand('clipboard')
        interp.call('rename', 'g12_original_clipboard', 'clipboard')
    assert copied[-1][-1] == expected
    target = tmp_path/'混合报告.txt'
    viewer.save(target)
    assert target.read_text(encoding='utf-8') == expected
    error = Mock()
    monkeypatch.setattr('txt_output.messagebox.showerror', error)
    viewer.save(tmp_path/'不存在'/'报告.txt')
    error.assert_called_once()
    assert viewer.T.get('1.0', 'end-1c') == expected
    blocks = viewer.table_blocks.copy()
    viewer.clear()
    assert not viewer.table_blocks and viewer.T.get('1.0', 'end-1c') == ''
    assert all(not block.winfo_exists() for block in blocks)
    Minijmp.print(app, tables[0])
    close = viewer.main.protocol('WM_DELETE_WINDOW')
    viewer.main.tk.call(close)
    assert app.table.tOut is None
    Minijmp.print(app, tables[-1])
    try:
        app.tOut.main.update()
        assert app.tOut is not viewer
        assert len(app.tOut.table_blocks) == 1
        assert app.tOut.table_blocks[0].table is tables[-1]
    finally:
        app.tOut.close()

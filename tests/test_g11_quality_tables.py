"""G11：全部质量报告生产入口、原 TXT、原精度及真实文件导出。"""
import ast
import math
from pathlib import Path
from types import ModuleType, SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PIL import Image
from source_baselines import read_source_baseline, normalize_p1_numeric, normalize_p2_legacy

import GRR
import cpk
import control_chart as spc
from MiniJMP import Minijmp
from stat_table import StatTable, export_excel, export_png
from test_g1_text_output import viewer
from test_g4_quality_reports import Y, PART, OP, NEST_Y, NEST_PART, NEST_OP, equal_result
from test_g4_quality_reports import equal_capability_result


ROOT = Path(__file__).resolve().parents[1]
BASE = '63f2191554176452ea544826ced714bef21f918c'
GRR_CASES = [(method, op, inter) for method in ('aiag', 'emp', 'xbar')
             for op in (False, True) for inter in ('Auto', 'On', 'Off')]
CASES = ([('grr', *case) for case in GRR_CASES]
         + [('direct', method, op, inter) for method in ('aiag', 'emp')
            for op in (False, True) for inter in (False, True)]
         + [('xbar_frame',)]
         + [('nested', hist) for hist in (None, 1., 20.)]
         + [('cpk', within, shape, ranges) for within in (False, True)
            for shape in ('individual', 'balanced', 'unbalanced') for ranges in (False, True)]
         + [('indices', name) for name in ('', 'Within', 'Overall (AKA Ppk in JMP)')]
         + [('nonconformance', within) for within in (False, True)]
         + [('spc', kind) for kind in ('IMR', 'I', 'MR', 'R', 'S', 'P', 'NP', 'C', 'U')])


@pytest.fixture(scope='module')
def baseline():
    modules = {}
    for name in ('GRR', 'cpk', 'control_chart'):
        source = read_source_baseline(BASE, f'code/{name}.py')
        module = ModuleType('g11_before_' + name)
        exec(compile(source, f'<G11-before:{name}>', 'exec'), module.__dict__)
        modules[name] = module
    return SimpleNamespace(**modules)


def report(case, grr=GRR, capability=cpk, control=spc):
    lines = []
    options = dict(print_port=lines.append, print_out=True)
    if case[0] in ('grr', 'direct'):
        _, method, has_op, inter = case
        options.update(y=Y, part=PART, op=OP if has_op else None)
        if case[0] == 'grr':
            result = grr.grr_master(mode=method, inter=inter, hist_std=30., spec_gap=80., **options)
        else:
            result = getattr(grr, 'grr_anova_' + method.upper())(inter=inter, **options)
    elif case[0] == 'xbar_frame':
        result = grr.GRR_xbar(pd.DataFrame({'o': OP, 'p': PART, 'v': Y}),
                              op_key='o', pt_key='p', y_key='v', **options)
    elif case[0] == 'nested':
        result = grr.grr_nested(y=NEST_Y, part=NEST_PART, op=NEST_OP, hist_std=case[1], **options)
    elif case[0] == 'cpk':
        _, within, shape, ranges = case
        data = {'individual': [1., 2., 4., 5., 7., 8.],
                'balanced': [[1., 2., 4.], [5., 7., 8.]],
                'unbalanced': [[1., 2., 4.], [5., 7.]]}[shape]
        result = capability.cpk(data, 4., 9., 0., within=within, use_range=ranges, **options)
    elif case[0] == 'indices':
        result = capability.cpk_calc(2., 1.3, 4., 9., 0., 19, 20, name=case[1], **options)
    elif case[0] == 'nonconformance':
        result = capability.nonconformance([-1., 1., 3., 5., 10.], 4., 0., 9., 3., 1.3, 2.,
                                           within=case[1], **options)
    else:
        kind = case[1]
        fig = Figure()
        ax = fig.subplots()
        options.update(fig=fig, ax=ax, xLabels=['甲', '乙', '丙', '丁'])
        if kind in ('IMR', 'I', 'MR', 'R', 'S'):
            data = [[1., 3.], [2., 5., 6.], [3., 4.], [2., 7.]] if kind in ('R', 'S') else [1., 3., 2., 5.]
            result = control.x_plot(data, plot_type='XBAR_'+kind if kind in ('R', 'S') else 'x_mR',
                                    one_chart=kind == 'I', chart_two_only=kind == 'MR', **options)
        elif kind in ('P', 'NP'):
            result = control.p_np_plot([.1, .2, .05, .15], n=[10, 20, 40, 30], NP=kind == 'NP', **options)
        else:
            result = control.cu_plot([1, 4, 2, 3], n=1 if kind == 'C' else [10, 20, 40, 30], C=kind == 'C', **options)
    return result, lines


@pytest.mark.parametrize('case', CASES, ids=str)
def test_all_report_producers_keep_txt_results_and_export_real_files(case, baseline, tmp_path, monkeypatch):
    result, lines = report(case)
    rendered = []
    module = (baseline.control_chart if case[0] == 'spc' else baseline.cpk
              if case[0] in ('cpk', 'indices', 'nonconformance') else baseline.GRR)
    original = module.gen_table
    def capture(rows):
        table = original(rows)
        rendered.append(table)
        return table
    monkeypatch.setattr(module, 'gen_table', capture)
    before, old_lines = report(case, baseline.GRR, baseline.cpk, baseline.control_chart)
    if case[0] in ('cpk', 'indices'):
        equal_capability_result(result, before)
    else:
        equal_result(result, before)
    assert '\n'.join(lines) == '\n'.join(old_lines)
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert tables
    assert len(tables) == sum(line.startswith('+') for line in old_lines)
    for index, table in enumerate(tables):
        assert table.headers == rendered[index].field_names
        assert table.display == [[str(cell) for cell in row] for row in rendered[index].rows]
        assert len(table.values) == len(table.display)
        assert all(len(row) == len(table.headers) for row in table.values)
        excel = tmp_path / f'{index}.xlsx'
        png = tmp_path / f'{index}.png'
        export_excel(table, excel)
        export_png(table, png)
        book = load_workbook(excel)
        sheet = book.active
        assert sheet.max_row == len(table.values) + 2
        assert sheet.max_column == len(table.headers)
        assert [c.value for c in sheet[2]] == table.headers
        for r, (values, shown) in enumerate(zip(table.values, table.display), 3):
            for c, (value, text) in enumerate(zip(values, shown), 1):
                cell = sheet.cell(r, c)
                if isinstance(value, (int, float, np.number)) and math.isfinite(value):
                    assert cell.data_type == 'n'
                    assert cell.value == pytest.approx(value, rel=2e-15, abs=1e-15)
                    assert cell.number_format == table.number_formats[c-1]
                else:
                    assert cell.value == text
        book.close()
        with Image.open(png) as image:
            assert image.format == 'PNG'
            assert image.height == 72 + 36 * (len(table.values) + 1)
            assert image.crop((0, image.height-45, image.width, image.height-10)).getextrema() != ((255, 255),)*3


@pytest.mark.parametrize('case', [c for c in CASES if c[0] != 'spc'], ids=str)
def test_raw_values_match_unrounded_baseline_expressions(case, baseline, monkeypatch):
    captured = []
    name = 'cpk' if case[0] in ('cpk', 'indices', 'nonconformance') else 'GRR'
    module = getattr(baseline, name)
    original = module.gen_table
    def capture(rows):
        captured.append(rows[1:])
        return original(rows)
    # 说明中的字符串拼接仍使用原 format_stat，仅将表格构造表达式解包。
    source = read_source_baseline(BASE, f'code/{name}.py')
    tree = ast.parse(source)
    class UnroundRows(ast.NodeTransformer):
        def visit_Assign(self, node):
            if any(isinstance(t, ast.Name) and t.id in ('rows', 'row', 'table', 't') for t in node.targets):
                node.value = Unround().visit(node.value)
            return self.generic_visit(node)
        def visit_Expr(self, node):
            if isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Attribute) and isinstance(node.value.func.value, ast.Name) and node.value.func.value.id in ('rows', 'row'):
                node.value = Unround().visit(node.value)
            return self.generic_visit(node)
        def visit_AugAssign(self, node):
            target = node.target.value if isinstance(node.target, ast.Subscript) else node.target
            if isinstance(target, ast.Name) and target.id in ('rows', 'row'):
                node.value = Unround().visit(node.value)
            return self.generic_visit(node)
    class Unround(ast.NodeTransformer):
        def visit_Call(self, node):
            node = self.generic_visit(node)
            if isinstance(node.func, ast.Name) and node.func.id == 'format_stat':
                return node.args[0]
            if isinstance(node.func, ast.Name) and node.func.id == 'str' and ast.unparse(node.args[0]) == 'DF':
                return node.args[0]
            return node
    namespace = dict(module.__dict__)
    from utilities import format_stat
    namespace['format_stat'] = format_stat
    exec(compile(ast.fix_missing_locations(UnroundRows().visit(tree)), '<unrounded-report>', 'exec'), namespace)
    namespace['gen_table'] = capture
    raw_module = SimpleNamespace(**namespace)
    report(case, grr=raw_module if name == 'GRR' else GRR,
           capability=raw_module if name == 'cpk' else cpk)
    _, lines = report(case)
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert len(tables) == len(captured)
    for table, expected in zip(tables, captured):
        if case[0] in ('cpk', 'indices') and table.headers == ['指标', '估计值', '置信下限', '置信上限']:
            for row, old in zip(table.values, expected):
                if row[0] in ('Cpl', 'Cpu', 'Ppl', 'Ppu'):
                    equal_result(row[:2], old[:2])
                    for a, b in zip(row[2:], old[2:]):
                        assert type(a) is type(b)
                        assert a == pytest.approx(b, rel=1e-11, abs=1e-12)
                else:
                    equal_result(row, old)
            assert len(table.values) == len(expected)
        else:
            equal_result(table.values, expected)


def test_spc_raw_limits_sizes_labels_missing_and_precision():
    lines = []
    spc._print_spc_report('U', [('U 图', [.123456789012345, np.nan, .2],
                              [0.], .123456789012345, [.3, .4])],
                          [10, 20, 40], '检查单位数', ['甲', '=标记', '丙'], ['原有警告'], lines.append)
    summary, detail = [line for line in lines if isinstance(line, StatTable)]
    assert summary.values == [['U 图', 0., .123456789012345, '逐点变化（见明细）']]
    assert detail.values[0] == [1, '甲', 10, .123456789012345, 0., .123456789012345, .3]
    assert math.isnan(detail.values[1][3])
    assert detail.values[-1][-1] is None
    assert detail.values[1][1] == '=标记'
    assert '原有警告' in lines
    assert all(not isinstance(line, StatTable) for line in lines if '未执行特殊原因' in line)


@pytest.mark.parametrize('kind', ['IMR', 'I', 'MR', 'R', 'S', 'P', 'NP', 'C', 'U'])
def test_every_spc_type_exports_exact_chart_values_and_limits(kind, baseline, monkeypatch):
    charts = []
    original = baseline.control_chart._print_spc_report
    def capture(title, panels, sizes, size_label, labels, notes, print_port):
        charts.append((panels, sizes, labels))
        original(title, panels, sizes, size_label, labels, notes, print_port)
    monkeypatch.setattr(baseline.control_chart, '_print_spc_report', capture)
    report(('spc', kind), control=baseline.control_chart)
    _, lines = report(('spc', kind))
    tables = [line for line in lines if isinstance(line, StatTable)]
    panels, sizes, labels = charts[0]
    assert len(tables) == len(panels) + 1
    def at(value, i):
        return value[i] if isinstance(value, (list, tuple, np.ndarray)) else value
    for table, (name, points, lower, center, upper) in zip(tables[1:], panels):
        assert table.title == name
        assert len(table.values) == len(points)
        for i, row in enumerate(table.values):
            expected = [i+1, labels[i], at(sizes, i), points[i],
                        at(lower, i), at(center, i), at(upper, i)]
            equal_result(row, expected)


@pytest.mark.parametrize('case', [('grr', 'emp', True, 'Auto'), ('nested', None),
                                  ('cpk', True, 'unbalanced', True), ('spc', 'U')])
def test_real_application_route_gui_and_txt(viewer, tmp_path, case, monkeypatch):
    _, lines = report(case)
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    for line in lines:
        Minijmp.print(app, line)
    viewer.main.geometry('480x320')
    viewer.main.update()
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert [b.table for b in viewer.table_blocks] == tables
    assert all(b.menu.entrycget(0, 'label') == '导出 Excel' and b.menu.entrycget(1, 'label') == '导出 PNG' for b in viewer.table_blocks)
    target = tmp_path / '报告.txt'
    viewer.save(str(target))
    assert target.read_text(encoding='utf-8') == '\n'.join(lines) + '\n'
    block = viewer.table_blocks[-1]
    import stat_table
    for index, extension in enumerate(('.xlsx', '.png')):
        destination = tmp_path / ('菜单' + extension)
        monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kw: str(destination))
        block.menu.invoke(index)
        assert destination.exists()
    assert viewer.ctrlfr.winfo_ismapped()
    viewer.clear()
    assert not viewer.table_blocks and viewer.T.get('1.0', 'end-1c') == ''


def test_plot_dialog_and_statistical_dependencies_unchanged():
    for name in ('qc_dialog.py', 'ctrlchart_dialog.py', 'cpk_plot.py', 'grr_plot.py',
                 'chart_style.py', 'utilities.py', 'two_way_anova.py', 'one_way_ANOVA.py'):
        # 合并时 main 已含 G9/G10 呈现迁移；仍逐字保护这三个共享模块。
        baseline = 'fc9215e0b7bb58d560754ff65607b2241d7ff7dc' if name in (
            'utilities.py', 'two_way_anova.py', 'one_way_ANOVA.py') else BASE
        before = read_source_baseline(baseline, f'code/{name}')
        current = (ROOT/'code'/name).read_text(encoding='utf-8')
        if name in ('qc_dialog.py', 'ctrlchart_dialog.py'):
            from source_baselines import normalize_column_memory
            current_tree = normalize_column_memory(ast.parse(current), name)
            if name == 'qc_dialog.py':
                normalize_p2_legacy(current_tree, name, ast.parse(before))
            assert ast.dump(current_tree) == ast.dump(ast.parse(before))
        elif name == 'utilities.py':
            # G12 仅迁移区间报告；其余全文及该函数计算仍受保护。
            trees = [ast.parse(source) for source in (before, current)]
            normalize_p1_numeric(trees[1], 'utilities.py', trees[0])
            for tree in trees:
                # G17 数字清洗由冻结输入对照及热点测试保护，其余源码继续比较。
                tree.body = [n for n in tree.body if not (
                    isinstance(n, ast.FunctionDef) and n.name == 'number_2Dlist')]
                function = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                                and n.name == 'mean_std_CIs')
                function.body = [n for n in function.body if not (
                    isinstance(n, ast.If) and isinstance(n.test, ast.Name) and n.test.id == 'print_out')]
            assert ast.dump(trees[0]) == ast.dump(trees[1])
        elif name == 'two_way_anova.py':
            trees = [ast.parse(source) for source in (before, current)]
            from source_baselines import normalize_review003
            normalize_review003(trees[1])
            for tree in trees:
                function = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                                and n.name == 'two_way_anova')
                branch = next(n for n in function.body if isinstance(n, ast.If)
                              and isinstance(n.test, ast.Name) and n.test.id == 'interaction')
                # G16 verifies the optimized cell SSE numerically and counts cell means.
                branch.body = branch.body[:1]
            assert ast.dump(trees[0]) == ast.dump(trees[1])
        else:
            assert current == before.replace('\r\n', '\n')


def test_dependency_freeze_rejects_unapproved_calculation(monkeypatch):
    read = Path.read_text
    target = ROOT / 'code' / 'qc_dialog.py'
    source = read(target, encoding='utf-8')
    assert source.count('P=P,') == 1
    monkeypatch.setattr(Path, 'read_text', lambda path, *a, **kw:
                        source.replace('P=P,', 'P=0.5,') if path == target else read(path, *a, **kw))
    with pytest.raises(AssertionError):
        test_plot_dialog_and_statistical_dependencies_unchanged()


@pytest.mark.parametrize('study', [False, True])
def test_grr_zero_variance_and_zero_optional_denominators(study):
    lines = []
    GRR._grr_tables([('总变异', 0.)], 0., lines.append, hist_std=0., spec_gap=0., study=study)
    variance, variation = [line for line in lines if isinstance(line, StatTable)]
    assert variance.values == [['总变异', 0., None]]
    assert variance.display == [['总变异', '0.000', '*']]
    assert variation.values == ([['总变异', 0., 0., None, None, None]] if study else [['总变异', 0.]])


@pytest.mark.parametrize('kind', ['Excel', 'PNG'])
def test_export_cancel_and_failure_stay_with_shared_component(viewer, monkeypatch, tmp_path, kind):
    import stat_table
    from unittest.mock import Mock
    viewer.add_table(StatTable(title='指标', headers=['值'], values=[[1.]]))
    block = viewer.table_blocks[0]
    exporter = Mock(side_effect=OSError('磁盘不可写'))
    monkeypatch.setattr(stat_table, 'export_excel' if kind == 'Excel' else 'export_png', exporter)
    error = Mock()
    monkeypatch.setattr(stat_table.messagebox, 'showerror', error)
    monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kw: '')
    block.export(kind)
    exporter.assert_not_called()
    error.assert_not_called()
    monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kw: str(tmp_path/'结果'))
    block.export(kind)
    exporter.assert_called_once()
    error.assert_called_once_with('导出失败', '磁盘不可写', parent=block)

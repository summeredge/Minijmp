"""G9：实际报告入口、固定 TXT、原始数值及独立文件导出。"""
import ast
import hashlib
import importlib
import json
import math
from numbers import Real
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from PIL import Image
from source_baselines import read_source_baseline, normalize_review003

import anova_dialog
import chi2_dialog
import qc_dialog
import sample_test_dialog
from MiniJMP import Minijmp
from stat_table import StatTable, export_excel, export_png
from test_g1_text_output import viewer

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / 'tests/fixtures/g9_reports.json'
MODULES = ['t_test', 'proportion_test', 'chi_sq', 'chi2_dialog',
           'one_way_ANOVA', 'two_way_anova', 'ti']
DATA = [1, 2, 3, 5, 7, 9]
OTHER = [2, 4, 5, 4, 8, 12]
GROUPS = [DATA, OTHER, [4, 6, 8, 11, 14, 17]]
LABELS = ['组甲', '组乙', '长标签丙']
CROSSED = [[[1, 2, 4], [3, 5, 6]], [[2, 5, 7], [6, 8, 11]]]
DIRECT = {
    't': ('t_test', 't_test_1sample', [DATA, 3], {}),
    't_summary': ('t_test', 't_test_1sample', [[], 3], dict(sample_n=12, sample_mean=4.123456, sample_std=2.34567, df=9.5)),
    'z': ('t_test', 'z_test_1sample', [DATA, 3, 2], {}),
    'z_summary': ('t_test', 'z_test_1sample', [[], 3, 2], dict(sample_n=12, sample_mean=4.123456)),
    'paired': ('t_test', 'paired_t_test', [DATA, OTHER[:-1]], {}),
    'welch': ('t_test', 't_test_2samples', [DATA, OTHER], {}),
    'pooled': ('t_test', 't_test_2samples', [DATA, OTHER], dict(pooled=True)),
    'two_t_summary': ('t_test', 't_test_2samples', [], dict(n1=12, n2=13, u1=3.123456, u2=5.789123, s1=2.123456, s2=4.987654)),
    'prop1': ('proportion_test', 'prop_test_1sample', [9, 40, .3], {}),
    'prop_zero': ('proportion_test', 'prop_test_1sample', [0, 12, .3], {}),
    'prop_all': ('proportion_test', 'prop_test_1sample', [12, 12, .3], {}),
    'prop2': ('proportion_test', 'prop_2sample', [], dict(n1=40, n2=50, e1=9, e2=18)),
    'sd_legacy': ('chi_sq', 'chi2_test_stdev', [DATA, 2], {}),
    'sd_alias': ('chi_sq', 'single_pop_var_test', [DATA, 2], {}),
    'sd_chi': ('chi_sq', 'chi2_var_1sample', [DATA], dict(s0=2)),
    'sd_chi_summary': ('chi_sq', 'chi2_var_1sample', [], dict(sample_n=12, sample_std=2.34567, s0=2)),
    'sd': ('chi_sq', 'var_1sample', [DATA], dict(s0=2)),
    'sd_summary': ('chi_sq', 'var_1sample', [], dict(sample_n=12, sample_std=2.34567, s0=2)),
    'multi_sd': ('chi_sq', 'multi_pop_var_test', [GROUPS, LABELS], {}),
    'two_sd': ('chi_sq', 'multi_pop_var_test', [GROUPS[:2], LABELS[:2]], {}),
    'chi': ('chi_sq', 'chi_square', [[[95, 43], [101, 64]]], {}),
    'contingency': ('chi_sq', 'chi_square', [[[12345678, 0], [9, 2]]], dict(print_table=True, rKeys=['甲', '乙'], cKeys=['中文列', '长标签列'])),
    'chi_missing': ('chi_sq', 'chi_square', [[[0, 0], [5, 5]]], {}),
    'one': ('one_way_ANOVA', 'one_way_anova', [GROUPS, LABELS], {}),
    'pairwise': ('one_way_ANOVA', 'JMP_ANOVA_t_test', [GROUPS, LABELS], {}),
    'two_inter': ('two_way_anova', 'two_way_anova', [CROSSED], dict(name_a='温度', name_b='批次')),
    'two_additive': ('two_way_anova', 'two_way_anova', [CROSSED], dict(interaction=False, name_a='温度', name_b='批次')),
    'ti': ('ti', 'ti_norm', [DATA], {}),
    'ti_summary': ('ti', 'ti_norm', [], dict(sample_n=10, sample_mean=5.95, sample_std=3.052, P=.85, alpha=.1)),
}
SAMPLE_DIALOGS = ['Mean1SampleDialog', 'Mean1SampleZDialog', 'Var1SampleDialog',
                  'Mean2SampleDialog', 'Prop1SDialog', 'Prop2SDialog']
DIALOGS = [f'{name}:{mode}' for name in SAMPLE_DIALOGS for mode in ('raw', 'summary')]
DIALOGS += ['PairedTDialog', 'MultiVarDialog', 'Anova1WayDialog_old', 'TtestDialog',
            'anova_columns', 'anova_stacked', 'anova_two_inter', 'anova_two_additive',
            'chi_raw', 'chi_summary', 'gof_raw', 'gof_summary', 'gof_normalized',
            'ti_dialog_raw', 'ti_dialog_summary']
CASES = list(DIRECT) + DIALOGS


def var(value):
    return SimpleNamespace(get=lambda: value)


def report(case):
    lines = []
    if case in DIRECT:
        module, name, args, kwargs = DIRECT[case]
        with np.errstate(divide='ignore', invalid='ignore'):
            getattr(importlib.import_module(module), name)(*args, **kwargs, print_out=True, print_port=lines.append)
        return lines
    fields = dict(xvar='X', yvar='Y', u0=3, s0=2, d0=1, alpha=.05, pooled=False,
        sample_mean='', sample_std='', sample_n='', u1='', s1='', n1='', u2='', s2='', n2='',
        event_name='1', event_name2='2', p0=.4, event_cnt='', ttl_cnt='', event_cnt2='', ttl_cnt2='',
        grouping='组别', comp_mean=True, show_plot=False, zvar='测量', inter=True, rvar='', cnt='', prop=.9)
    if case.endswith('summary'):
        fields.update(sample_mean=30, sample_std=4, sample_n=20, u1=30, s1=4, n1=20,
                      u2=35, s2=6, n2=25, event_cnt=3, ttl_cnt=20, event_cnt2=8, ttl_cnt2=25)
    dialog = SimpleNamespace(remember_columns=Mock(), **{key: var(value) for key, value in fields.items()},
        df=pd.DataFrame({'X': DATA, 'Y': OTHER}), grpcols=['X', 'Y'],
        app=SimpleNamespace(print=lines.append), error=Mock(), parent=None)
    if ':' in case:
        getattr(sample_test_dialog, case.split(':')[0]).apply(dialog)
    elif case in ('PairedTDialog', 'MultiVarDialog'):
        getattr(sample_test_dialog, case).apply(dialog)
    elif case in ('Anova1WayDialog_old', 'TtestDialog', 'anova_columns'):
        getattr(anova_dialog, 'Anova1WayDialog' if case == 'anova_columns' else case).apply(dialog)
    elif case == 'anova_stacked':
        dialog.df = pd.DataFrame({'X': DATA + OTHER, '组别': ['甲'] * 6 + ['乙'] * 6})
        dialog.grpcols = []
        dialog.make_working_df = lambda **kw: anova_dialog.Anova1WayDialog.make_working_df(dialog, **kw)
        anova_dialog.Anova1WayDialog.apply(dialog)
    elif case.startswith('anova_two'):
        dialog.df = pd.DataFrame({'X': ['低']*6 + ['高']*6, 'Y': (['甲']*3 + ['乙']*3)*2,
                                 '测量': [v for a in CROSSED for b in a for v in b]})
        dialog.inter = var(case.endswith('inter'))
        anova_dialog.Anova2WayDialog.apply(dialog)
    elif case.startswith('chi_'):
        dialog.df = pd.DataFrame({'X': ['甲', '甲', '乙', '乙'], 'Y': ['A', 'B', 'A', 'B']})
        dialog.grpcols = []
        if case == 'chi_summary':
            dialog.df = pd.DataFrame({'X': [95, 101], 'Y': [43, 64], '标签': ['甲', '乙']})
            dialog.grpcols, dialog.rvar = ['X', 'Y'], var('标签')
        chi2_dialog.Chi2TableDialog.apply(dialog)
    elif case.startswith('gof_'):
        dialog.df = pd.DataFrame({'X': ['甲', '甲', '乙', '乙', '乙', '乙']})
        if case != 'gof_raw':
            dialog.df = pd.DataFrame({'X': ['甲', '乙'], '计数': [20, 30]})
            dialog.rvar, dialog.cnt = var('X'), var('计数')
        probabilities = [2., 3.] if case == 'gof_normalized' else [.4, .6]
        with patch.object(chi2_dialog, 'ProbabilityDialog', return_value=SimpleNamespace(result=True, results=probabilities)):
            chi2_dialog.Chi2PropDialog.apply(dialog)
    elif case.startswith('ti_dialog'):
        qc_dialog.TIDialog.apply(dialog)
    else:
        raise AssertionError(case)
    assert dialog.error.call_count == int(case == 'gof_normalized')
    return lines


class ComputationOnly(ast.NodeTransformer):
    def visit_If(self, node):
        if isinstance(node.test, ast.Name) and node.test.id == 'print_out':
            return None
        return self.generic_visit(node)

    def visit_ImportFrom(self, node):
        return None if node.module == 'stat_table' else node


def computation_hash(module):
    tree = ast.parse((ROOT / 'code' / f'{module}.py').read_bytes())
    if module == 'two_way_anova':
        normalize_review003(tree)
    if module == 'proportion_test':
        # G20 LR has an independent oracle; freeze every other computation as before.
        frozen = ast.parse(read_source_baseline(
            '11c4a7569f5b80836de04f06e4049c99ed1d084c', 'code/proportion_test.py'))
        old_lr = next(n for n in frozen.body if isinstance(n, ast.FunctionDef)
                      and n.name == 'p_likelihood_ratio')
        tree.body = [old_lr if isinstance(n, ast.FunctionDef) and n.name == 'p_likelihood_ratio'
                     else n for n in tree.body
                     if not (isinstance(n, ast.ImportFrom) and n.module == 'scipy.special'
                             and [(a.name, a.asname) for a in n.names] == [('rel_entr', None)])]
        # REVIEW-004/006/007 仅规范化下列精确语句；独立参考见 test_review_proportions.py。
        # 其余语句及函数仍参与原哈希，不能借授权入口放行新的计算差异。
        old_prop = next(n for n in frozen.body if isinstance(n, ast.FunctionDef)
                        and n.name == 'prop_2sample')
        start = next(i for i, n in enumerate(old_prop.body)
                     if isinstance(n, ast.Assign) and ast.unparse(n.targets[0]) == 'mode')
        changes = [
            (None, 'from scipy.stats import norm, binom, hypergeom, fisher_exact',
             ast.parse('from scipy.stats import norm, binom, hypergeom').body),
            ('prop_2sample', 'p_f = fisher_exact([[e1, n1 - e1], [e2, n2 - e2]]).pvalue',
             old_prop.body[start:start + 3]),
            ('prop_2sample', 'if SE_sep == 0:\n    CI = CI2 = (np.nan, np.nan)', []),
            ('prop_2sample', 'Zp = pd / SE_pooled if SE_pooled > 0 else np.nan',
             ast.parse('Zp = pd / SE_pooled').body),
            ('prop_2sample', 'Zs = pd / SE_sep if SE_sep > 0 else np.nan',
             ast.parse('Zs = pd / SE_sep').body),
        ]
        for name, endpoint in [('p_exct_l', '0.0 if events == 0'),
                               ('p_exct_low', '0.0 if events == 0'),
                               ('p_exct_u', '1.0 if events == N'),
                               ('p_exct_high', '1.0 if events == N')]:
            formula = 'dfn * F / (dfd + dfn * F)'
            changes.append(('prop_test_1sample', f'{name} = {endpoint} else {formula}',
                            ast.parse(f'{name} = {formula}').body))
        for owner, known, original in changes:
            body = tree.body if owner is None else next(
                n.body for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == owner)
            expected = ast.dump(ast.parse(known).body[0])
            matches = [i for i, node in enumerate(body) if ast.dump(node) == expected]
            assert len(matches) == 1, f'Unexpected REVIEW proportion AST: {known}'
            body[matches[0]:matches[0] + 1] = original
    if module == 'chi2_dialog':
        # 此 Dialog 的报告没有 print_out 分支；由固定 TXT 和原有路由测试验证。
        return None
    return hashlib.sha256(ast.dump(ComputationOnly().visit(tree)).encode()).hexdigest()


@pytest.mark.parametrize('old,new', [
    ('[e2, n2 - e2]]).pvalue', "[e2, n2 - e2]], alternative='less').pvalue"),
    ('Zs = pd / SE_sep if SE_sep > 0 else np.nan',
     'Zs = pd / SE_pooled if SE_sep > 0 else np.nan'),
    ('0.0 if events == 0 else', '0.1 if events == 0 else'),
    ('if probability <= PTK', 'if probability < PTK'),
])
def test_review_proportion_normalization_rejects_other_changes(monkeypatch, old, new):
    path = ROOT / 'code/proportion_test.py'
    source = path.read_bytes()
    assert source.count(old.encode()) >= 1
    changed = source.replace(old.encode(), new.encode(), 1)
    read_bytes = Path.read_bytes
    monkeypatch.setattr(Path, 'read_bytes',
                        lambda self: changed if self == path else read_bytes(self))
    baseline = json.loads(FIXTURE.read_text(encoding='utf-8'))
    with pytest.raises(AssertionError):
        assert computation_hash('proportion_test') == baseline['computation']['proportion_test']


@pytest.mark.parametrize('case', ['prop_zero', 'prop_all'])
def test_cp_endpoint_fixture_matches_binomial_tail_inversion(case):
    # 独立推导：P(X=0)=(1-p)^12，P(X=12)=p^12；双侧尾概率 .025，单侧 .05。
    # 因而仅将 CP 不存在事件一侧的旧 NaN 文本改为边界 0/1，不重生成报告。
    if case == 'prop_zero':
        limits = [0, 1 - .025**(1/12), 0, 1 - .05**(1/12)]
    else:
        limits = [.025**(1/12), 1, .05**(1/12), 1]
    baseline = json.loads(FIXTURE.read_text(encoding='utf-8'))
    row = next(line for line in baseline['text'][case].splitlines()
               if line.startswith('| Clopper-Pearson |'))
    assert [cell.strip() for cell in row.split('|')[2:-1]] == [f'{x:.3f}' for x in limits]


@pytest.mark.parametrize('case', CASES)
def test_fixed_txt_and_every_table_reaches_viewer(viewer, tmp_path, case):
    baseline = json.loads(FIXTURE.read_text(encoding='utf-8'))
    lines = report(case)
    assert '\n'.join(lines) + '\n' == baseline['text'][case]
    tables = [line for line in lines if isinstance(line, StatTable)]
    assert len(tables) == baseline['tables'][case]
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    for line in lines:
        Minijmp.print(app, line)
    viewer.main.update()
    assert [block.table for block in viewer.table_blocks] == tables
    target = tmp_path / '原报告.txt'
    viewer.save(target)
    assert target.read_text(encoding='utf-8') == baseline['text'][case]
    assert all(block.menu.index('end') == 2 for block in viewer.table_blocks)


@pytest.mark.parametrize('case', list(DIRECT) + ['gof_summary'])
def test_each_table_real_excel_png_complete(tmp_path, case):
    tables = [line for line in report(case) if isinstance(line, StatTable)]
    assert tables
    for index, table in enumerate(tables):
        assert table.title and table.headers
        assert len(table.values) == len(table.display)
        xlsx = tmp_path / f'{case}-{index}.xlsx'
        png = xlsx.with_suffix('.png')
        export_excel(table, xlsx)
        book = load_workbook(xlsx)
        sheet = book.active
        assert sheet.cell(1, 1).value == table.title
        assert [cell.value for cell in sheet[2]] == table.headers
        assert sheet.max_row == len(table.values) + 2
        assert sheet.max_column == len(table.headers)
        for r, (values, display) in enumerate(zip(table.values, table.display), 3):
            for c, (value, shown) in enumerate(zip(values, display), 1):
                cell = sheet.cell(r, c)
                if isinstance(value, Real) and math.isfinite(value):
                    assert cell.data_type == 'n'
                    assert cell.value == pytest.approx(value, rel=1e-14, abs=0)
                    assert cell.number_format == table.number_formats[c-1]
                else:
                    assert cell.value == (shown or None)
        book.close()
        export_png(table, png)
        with Image.open(png) as image:
            assert image.format == 'PNG'
            assert image.height == 72 + 36 * (len(table.values) + 1)
            # 每行包含最后一行均有内容；G8 已覆盖逐列绘制与宽表裁切。
            for row in range(len(table.values) + 1):
                assert image.crop((0, 60+36*row, image.width, 90+36*row)).getextrema() != ((255, 255),)*3


@pytest.mark.parametrize('module', MODULES)
def test_no_string_table_ports_and_computations_unchanged(module):
    source = (ROOT / 'code' / f'{module}.py').read_text(encoding='utf-8')
    tree = ast.parse(source)
    for node in ast.walk(tree):
        if isinstance(node, ast.Call) and ast.unparse(node.func) in ('print', 'print_port', 'self.app.print'):
            assert not (node.args and isinstance(node.args[0], ast.Call)
                        and ast.unparse(node.args[0].func) == 'str'), node.lineno
    baseline = json.loads(FIXTURE.read_text(encoding='utf-8'))
    assert computation_hash(module) == baseline['computation'][module]


@pytest.mark.parametrize('case,checks', [
    ('t', [(1, 0, 2, ('t',)), (1, 0, 3, ('df',)), (1, 0, 4, ('p',)),
           (2, 0, 1, ('u0_95range', 0)), (2, 0, 2, ('u0_95range', 1))]),
    ('z', [(1, 0, 2, ('z',)), (1, 0, 4, ('p',))]),
    ('paired', [(1, 0, 2, ('t',)), (1, 0, 4, ('p',))]),
    ('welch', [(1, 0, 3, ('df',)), (1, 0, 4, ('p',)), (2, 0, 1, ('CI', 0))]),
    ('pooled', [(1, 0, 3, ('df',)), (1, 0, 4, ('p',)), (2, 0, 2, ('CI', 1))]),
    ('prop1', [(1, 0, 1, ('p norm',)), (1, 1, 1, ('p Sterne',)),
               (1, 2, 1, ('p Blaker',)), (1, 3, 1, ('Blaker p',)),
               (2, 0, 1, ('CI norm', 'CI l')), (2, 1, 2, ('CI exact', 'CI u')),
               (2, 2, 1, ('Blaker CI', 0)), (3, 0, 1, ('JMP chi sq', 'chi sq Pearson'))]),
    ('prop2', [(1, 0, 2, ('p Fisher',)), (1, 1, 2, ('p sep',)), (1, 2, 2, ('p Pooled',))]),
    ('sd_legacy', [(1, 0, 2, ('chi_sq',)), (1, 0, 4, ('p',)),
                   (2, 0, 1, ('95_range_from_S', 0)), (3, 0, 1, ('95_chi2_range', 0))]),
    ('sd_chi', [(1, 0, 4, ('p',)), (2, 0, 2, ('CI 95', 1))]),
    ('sd', [(1, 0, 1, ('p',)), (1, 1, 1, ('Bonett p',)),
            (2, 1, 2, ('Bonett CI', 1))]),
    ('chi', [(0, 0, 1, ('chi sq Pearson',)), (0, 1, 3, ('p LR',))]),
    ('one', [(0, 0, 4, ('CI', 0, 0)), (1, 0, 3, ('MSB',)),
             (1, 0, 4, ('F',)), (1, 0, 5, ('p',)), (1, 1, 3, ('MSW',))]),
    ('pairwise', [(0, 0, 2, ('p', 1, 0)), (0, 2, 2, ('p', 2, 1))]),
    ('two_inter', [(1, 0, 2, ('SSA',)), (1, 0, 5, ('p A',)),
                   (1, 1, 2, ('SSB',)), (1, 2, 2, ('SSAB',)), (1, 3, 3, ('MSE',))]),
    ('two_additive', [(1, 0, 2, ('SSA',)), (1, 1, 5, ('p B',)), (1, 2, 3, ('MSE',))]),
    ('ti_summary', [(1, 0, 4, ('TI', 0)), (1, 0, 5, ('TI', 1)),
                    (1, 1, 4, ('Lower Bound',)), (1, 2, 5, ('Upper Bound',))]),
])
def test_export_values_keep_exact_computation_precision(case, checks):
    module, name, args, kwargs = DIRECT[case]
    lines = []
    result = getattr(importlib.import_module(module), name)(
        *args, **kwargs, print_out=True, print_port=lines.append)
    tables = [line for line in lines if isinstance(line, StatTable)]
    for table, row, col, path in checks:
        expected = result
        for key in path:
            expected = expected[key]
        assert tables[table].values[row][col] == expected


def test_nonstandard_layout_keeps_both_numbers_and_percentages():
    from scipy.stats import chi2_contingency
    from chi_sq import chi2_test_stdev
    contingency = next(line for line in report('contingency') if isinstance(line, StatTable))
    counts = [[12345678, 0], [9, 2]]
    expected = chi2_contingency(counts, correction=False).expected_freq
    assert contingency.headers == ['行 / 列', '中文列（观测）', '中文列（期望）',
                                   '长标签列（观测）', '长标签列（期望）', '合计']
    for row in range(2):
        assert contingency.values[row][1::2] == counts[row] + [sum(counts[row])]
        assert contingency.values[row][2:5:2] == list(expected[row])
    assert contingency.values[-1] == ['合计', 12345687, '*', 2, '*', 12345689]
    multi = next(line for line in report('multi_sd') if isinstance(line, StatTable))
    assert len(multi.headers) == 7
    for data, row in zip(GROUPS, multi.values):
        assert tuple(row[3:5]) == chi2_test_stdev(data, 1)['95_range_from_S']
    ti = [line for line in report('ti_summary') if isinstance(line, StatTable)][1]
    assert all(row[2:4] == [.85, .9] for row in ti.values)
    assert ti.number_formats[2:4] == ['0.0%', '0.0%']


def test_dynamic_numeric_labels_and_special_symbols_are_preserved():
    from chi_sq import chi_square
    lines = []
    chi_square([[0, 10], [5, 5]], print_port=lines.append, print_table=True,
               rKeys=[1, '乙 / μ'], cKeys=[2, '列 / σ'])
    table = next(line for line in lines if isinstance(line, StatTable))
    assert table.headers == ['行 / 列', '2（观测）', '2（期望）',
                             '列 / σ（观测）', '列 / σ（期望）', '合计']
    assert table.display[0][0] == '1' and table.display[1][0] == '乙 / μ'
    assert '2/4 个单元格的期望频数小于 5' in '\n'.join(lines)

"""G14：真实 StatTable → Minijmp.print → txt_viewer → TableBlock 链路。"""
import gc
import json
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from unittest.mock import Mock
import weakref
from tkinter import ttk

import pytest
from openpyxl import load_workbook
from PIL import Image
from prettytable import PrettyTable

from MiniJMP import Minijmp
import MiniJMP
import stat_table
from stat_table import StatTable, TableBlock
from test_g1_text_output import viewer


def assert_worker_memory(result):
    assert result['sampled_process_peak_rss_mib'] >= result['memory_imported_mib']['WorkingSetSize'], (
        '父进程没有采样到实际工作进程的内存')


@pytest.mark.skipif(os.name != 'nt', reason='Windows 进程内存采样')
@pytest.mark.parametrize('launcher_memory', [False, True])
def test_windows_benchmark_exports_smoke(tmp_path, launcher_memory):
    script = Path(__file__).resolve().parents[1] / 'tools/benchmark_report_rendering.py'
    output = tmp_path / '基准 smoke'
    arguments = ['--scenario', 'small', '--phase', 'exports', '--output', str(output),
                 '--export-timeout', '20']
    command = [sys.executable, str(script), *arguments]
    if launcher_memory:
        # 仅替换父进程的采样；真实 worker 仍在独立进程报告自己的 PID 和工作集。
        source = (
            "import runpy,sys; benchmark=runpy.run_path(sys.argv[1]); "
            "benchmark['main'].__globals__['memory']=lambda pid=None: {'PeakWorkingSetSize': 1.0}; "
            "sys.argv=sys.argv[1:]; benchmark['main']()")
        command = [sys.executable, '-c', source, str(script), *arguments]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=35)
    assert completed.returncode == 0, completed.stderr
    assert [path.name for path in output.iterdir()] == ['small-exports.json']
    result = json.loads((output / 'small-exports.json').read_text(encoding='utf-8'))
    assert result['status'] == 'ok' and result['stage'] == 'complete', result
    assert result['phase'] == 'exports' and result['shape'] == [100, 7] and result['tables'] == 1
    assert isinstance(result['worker_pid'], int) and result['worker_pid'] > 0
    assert result['worker_pid'] != os.getpid()
    assert Path(result['python']).resolve() == Path(sys.executable).resolve()
    if launcher_memory:
        with pytest.raises(AssertionError, match='实际工作进程'):
            assert_worker_memory(result)
    else:
        assert_worker_memory(result)


def report(rows=100, columns=7, **kwargs):
    return StatTable(title='公共渲染 μ', headers=['变量'] + [f'列{c}' for c in range(1, columns)],
                     values=[[f'样本{r+1}'] + [r+c/7 for c in range(1, columns)]
                             for r in range(rows)], **kwargs)


def emit(viewer, *items):
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    for item in items:
        Minijmp.print(app, item)
    viewer.main.update()
    return viewer.table_blocks[-1]


def widget_count(widget):
    return 1 + sum(widget_count(child) for child in widget.winfo_children())


def test_small_100_by_7_keeps_labels_and_text(viewer):
    table = report()
    block = emit(viewer, table)
    assert isinstance(block, TableBlock) and not block.paged
    labels = block.winfo_children()[1].winfo_children()
    assert len(labels) == 101 * 7 and all(isinstance(label, ttk.Label) for label in labels)
    assert [label.cget('text') for label in labels[-7:]] == table.display[-1]
    assert viewer.T.get('1.0', 'end-1c') == str(table)+'\n'


@pytest.mark.parametrize('rows,columns', [(1000, 7), (10000, 7), (50000, 7), (100, 100)])
def test_bounded_tree_reaches_last_row_and_column(viewer, rows, columns):
    table = report(rows, columns)
    block = emit(viewer, table)
    assert block.paged and widget_count(block) < 30
    assert len(block.tree.get_children()) <= block.page_size <= 100
    assert len(block.tree.get_children()) * columns <= 1000
    assert len(viewer.T.get('1.0', 'end-1c')) == 1  # 不向 Tk 再复制完整 ASCII 文本。
    block.show_page(block.page_count)
    viewer.main.update()
    last = block.tree.get_children()[-1]
    assert last == str(rows-1)
    assert list(block.tree.item(last, 'values')) == table.display[-1]
    assert str(block.tree.heading(str(columns-1), 'text')) == table.headers[-1]
    block.tree.xview_moveto(1)
    block.tree.yview_moveto(1)
    viewer.main.update()
    if columns == 100:
        assert block.tree.xview()[0] > 0
    assert block.tree.yview()[1] == 1
    block.show_page(1)
    assert list(block.tree.item('0', 'values')) == table.display[0]


def test_pages_keyboard_spinbox_scroll_and_no_loss(viewer):
    table = report(1003)
    block = emit(viewer, table)
    seen = []
    for page in range(1, block.page_count+1):
        block.show_page(page)
        block.tree.yview_moveto(1)
        viewer.main.update()
        seen.extend([list(block.tree.item(item, 'values')) for item in block.tree.get_children()])
    assert seen == table.display
    block.show_page(1)
    block.tree.focus_force()
    block.tree.event_generate('<Next>')
    viewer.main.update()
    assert block.page == 2
    block.tree.event_generate('<Control-End>')
    viewer.main.update()
    assert block.page == block.page_count
    block.page_number.set('2')
    block.page_input.focus_force()
    block.page_input.event_generate('<Return>')
    viewer.main.update()
    assert block.page == 2
    block.page_number.set('无效')
    block.page_input.event_generate('<Return>')
    viewer.main.update()
    assert block.page == 2 and block.page_number.get() == '2'
    block.show_page(99999)
    assert block.page == block.page_count
    block.show_page(-1)
    assert block.page == 1


def test_mixed_txt_save_and_copy_preserve_legacy_text_and_selection(viewer, tmp_path, monkeypatch):
    small = StatTable('原有小表文本', '小表', ['甲'], [[1]])
    large = report(1000)
    items = ['前文', small, large.title, large, '后文']
    emit(viewer, *items)
    expected = ''.join(str(item)+'\n' for item in items)
    target = tmp_path / '完整报告.txt'
    viewer.save(target)
    assert target.read_text(encoding='utf-8') == expected
    copied = []
    monkeypatch.setattr(viewer.T, 'clipboard_clear', lambda: None)
    monkeypatch.setattr(viewer.T, 'clipboard_append', copied.append)
    viewer.select_all()
    viewer.T.event_generate('<<Copy>>')
    assert copied == [expected]
    tree = viewer.table_blocks[-1].tree
    tree.focus_force()
    tree.event_generate('<Control-a>')
    viewer.main.update()
    tree.event_generate('<<Copy>>')
    assert copied[-1] == expected
    index = next(index for kind, name, index in viewer.T.dump('1.0', 'end', window=True)
                 if name == str(viewer.table_blocks[-1]))
    viewer.T.tag_remove('sel', '1.0', 'end')
    viewer.T.tag_add('sel', index, f'{index}+1c')
    viewer.T.event_generate('<<Copy>>')
    assert copied[-1] == str(large)
    viewer.T.tag_remove('sel', '1.0', 'end')
    viewer.T.tag_add('sel', '1.1', '1.2')
    viewer.T.event_generate('<<Copy>>')
    assert copied[-1] == '文'
    error = Mock()
    monkeypatch.setattr('txt_output.messagebox.showerror', error)
    viewer.save(tmp_path / '不存在' / '报告.txt')
    error.assert_called_once()
    viewer.save(target)
    assert target.read_text(encoding='utf-8') == expected


def test_large_excel_full_precision_formats_and_menu(viewer, tmp_path, monkeypatch):
    values = [['=标签', .123456789012345, r, None] for r in range(1003)]
    table = StatTable('原始报告', '完整导出', ['变量', '比例', 'N', '缺失'], values,
                      number_formats=['General', '0.0000%', '0', '0.000'])
    block = emit(viewer, table)
    block.show_page(block.page_count)
    target = tmp_path / '完整.xlsx'
    monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kw: str(target))
    block.menu.invoke(0)
    book = load_workbook(target)
    sheet = book.active
    assert sheet.max_row == 1005 and sheet.max_column == 4
    assert sheet['A1'].value == table.title and [c.value for c in sheet[2]] == table.headers
    for r in range(3, 1006):
        assert sheet.cell(r, 1).value == '=标签' and sheet.cell(r, 1).data_type == 's'
        assert sheet.cell(r, 2).value == .123456789012345
        assert sheet.cell(r, 2).number_format == '0.0000%'
        assert sheet.cell(r, 3).value == r-3 and sheet.cell(r, 3).number_format == '0'
        assert sheet.cell(r, 4).value == '*'
    assert sheet.freeze_panes == 'B3' and sheet['A1'].font.bold
    assert sheet['B3'].alignment.horizontal == 'right'
    book.close()


def test_small_png_complete_and_huge_png_guard_before_allocation(viewer, tmp_path, monkeypatch):
    small = report(3)
    block = emit(viewer, small)
    target = tmp_path / '完整.png'
    monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **kw: str(target))
    block.menu.invoke(1)
    with Image.open(target) as image:
        assert image.height == 72+36*4 and image.getbbox()
    large = report(1000)
    block = emit(viewer, large)
    allocate = Mock(side_effect=AssertionError('不得分配超大位图'))
    monkeypatch.setattr(Image, 'new', allocate)
    error = Mock()
    monkeypatch.setattr(stat_table.messagebox, 'showerror', error)
    block.menu.invoke(1)
    assert 'Excel' in error.call_args.args[1]
    assert 'PNG' in error.call_args.args[1]
    allocate.assert_not_called()


def test_wide_png_guard_before_allocation(tmp_path, monkeypatch):
    table = StatTable('兼容文本', '宽表', ['变量', '宽列'], [['行', 'W'*10000]])
    allocate = Mock(side_effect=AssertionError('不得分配超大位图'))
    monkeypatch.setattr(Image, 'new', allocate)
    with pytest.raises(ValueError, match='Excel'):
        stat_table.export_png(table, tmp_path / '宽.png')
    allocate.assert_not_called()


def test_png_pixel_budget_before_allocation(tmp_path, monkeypatch):
    table = report(100, 100)
    allocate = Mock(side_effect=AssertionError('不得分配超大位图'))
    monkeypatch.setattr(Image, 'new', allocate)
    with pytest.raises(ValueError, match='MiB.*Excel'):
        stat_table.export_png(table, tmp_path / '面积.png')
    allocate.assert_not_called()


def test_twenty_tables_budget_scroll_clear_close_reopen(viewer, monkeypatch):
    table = report()
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    initial = widget_count(viewer.main)
    for _ in range(20):
        Minijmp.print(app, table)
    assert len(viewer.table_blocks) == 20
    assert len(viewer.T.tk.call('after', 'info')) >= 1
    viewer.main.update()
    assert widget_count(viewer.main) < 2500
    for kind, name, index in viewer.T.dump('1.0', 'end', window=True):
        viewer.T.see(index)
        viewer.main.update()
        assert viewer.T.nametowidget(name).winfo_ismapped()
    assert ''.join(viewer._text_chunks()) == (str(table)+'\n')*20
    blocks = list(viewer.table_blocks)
    viewer.clear()
    assert viewer.table_blocks == [] and widget_count(viewer.main) == initial
    assert all(not block.winfo_exists() and block.table is None for block in blocks)
    Minijmp.print(app, table)
    block = weakref.ref(viewer.table_blocks[0])
    viewer.close()  # 包括尚未执行的滚动回调。
    gc.collect()
    assert not viewer.table_blocks and app.table.tOut is None and block() is None
    assert viewer._scroll_job is None
    factory = Mock(wraps=MiniJMP.txt_viewer)
    monkeypatch.setattr(MiniJMP, 'txt_viewer', factory)
    Minijmp.print(app, report(1000))
    try:
        app.tOut.main.update()
        factory.assert_called_once()
        assert app.tOut.table_blocks[0].paged
    finally:
        app.tOut.close()


def test_stat_table_string_legacy_generators_and_snapshot():
    original = [['中文', .123456789012345]]
    legacy = PrettyTable(['变量', '比例'])
    legacy.add_row(['中文', '12.35%'])
    table = StatTable(legacy, '旧调用', iter(['变量', '比例']), (r for r in original),
                      display=(r for r in [['中文', '12.35%']]),
                      number_formats=['General', '0.00%'])
    assert isinstance(table, str) and str(table) == str(legacy)
    original[0][1] = 0
    assert table.values == [['中文', .123456789012345]]
    assert table.display == [['中文', '12.35%']]
    baseline = Path(__file__).resolve().parents[1] / 'docs/g14-evidence/before/stat_table.py'
    namespace = {'__name__': 'g14_baseline'}
    exec(compile(baseline.read_bytes(), str(baseline), 'exec'), namespace)
    before = namespace['StatTable'](title='公共渲染 μ', headers=['变量', '列1'],
                                    values=[['样本1', 1/7]])
    assert report(1, 2) == before

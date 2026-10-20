"""G10 的固定迁移前文本、原值及真实导出契约。"""
import json
import math
import ast
from pathlib import Path

import numpy as np
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PIL import Image
from types import SimpleNamespace
from unittest.mock import Mock

import dist_fit
import fit_dialog
import linear_fit
import linear_plot
import nonpara_cor
import ortho_fit
import Residual
import utilities
from test_g1_text_output import viewer

X = [1., 2., 4., 5., 7., 8.]
Y = [2., 4., 3., 8., 6., 10.]
DATA = [1., 2., 3., 4., 2., 12., 3., 12., 3., 2., 40., 3., 2.]
FIXTURE = Path(__file__).parent / 'fixtures' / 'g10_before.json'


def reports():
    cases = {
        'ols': lambda p: linear_fit.linear_fit(X, Y, print_port=p),
        'lack': lambda p: linear_fit.linear_fit([1, 1, 2, 2, 3, 3], Y, print_port=p),
        'perfect': lambda p: linear_fit.linear_fit(X, [2*x+1 for x in X], print_port=p),
        'ortho': lambda p: ortho_fit.orthogonal_fit(X, Y, print_residual=True, print_port=p),
        'residual_only': lambda p: ortho_fit.orthogonal_fit(X, Y, print_out=False, print_residual=True, print_port=p),
        'pearson': lambda p: utilities.pearson_correlation(X, Y, print_port=p),
        'normality': lambda p: utilities.norm_test(Y, print_out=True, print_port=p),
        'resid': lambda p: Residual.resid_test(Y, X, axs=Figure().subplots(2, 2), print_out=True, print_port=p),
        'fitted_resid': lambda p: Residual.linear_fit_resid_test(X, Y, axs=Figure().subplots(2, 2), print_out=True, print_port=p),
        'compare': lambda p: dist_fit.compare_dist(DATA, print_port=p),
        'nonfinite': lambda p: dist_fit.print_table('特殊', ['小值', '大值', '缺失', '无穷'],
            [1e-80, 1e80, np.nan, np.inf], [np.nan, np.inf, -np.inf], .3, 1e-100, print_port=p),
    }
    for name in ('fit_norm', 'fit_t', 'fit_gamma', 'fit_lognorm', 'fit_expon', 'fit_weibull'):
        cases[name] = lambda p, name=name: getattr(dist_fit, name)(DATA, print_port=p)
    for name in ('spearman', 'kendall', 'hoeffding'):
        func = getattr(nonpara_cor, name)
        cases[name] = lambda p, func=func: func(X, Y, print_out=True, print_port=p)
        cases[name+'_matrix'] = lambda p, func=func, name=name: fit_dialog.print_multi_col_tables(
            [X, Y], ['温度', '压力'], func, name=name, print_port=p)
    for name in ('multi_fit', 'multi_fit_orig'):
        cases[name] = lambda p, name=name: getattr(linear_plot, name)(
            [X, Y], ['温度', '压力'], print_out=True, print_port=p,
            **({'ax': Figure().subplots(), 'fig': Figure()} if name == 'multi_fit' else
               {'axs': Figure().subplots(1, 1, squeeze=False), 'fig': Figure()}))
    return cases


@pytest.mark.parametrize('name', list(reports()))
def test_text_values_and_real_exports(name, tmp_path):
    from stat_table import StatTable, export_excel, export_png
    outputs = []
    result = reports()[name](outputs.append)
    assert '\n'.join(outputs) == json.loads(FIXTURE.read_text(encoding='utf-8'))[name]
    tables = [o for o in outputs if isinstance(o, StatTable)]
    assert tables, name
    if name == 'ols':
        assert tables[0].values[0] == [result[k] for k in
            ['N', 'R Square', 'R Square Adj', 'Root MSE', 'Mean of Response']]
        for row, keys in zip(tables[1].values, [
                ['intercept', 'Std Error Int.', 't int.', 'p int.'],
                ['slope', 'Std Error Slope', 't slope', 'p slope']]):
            assert row[1:] == [result[k] for k in keys]
    elif name in ('ortho', 'residual_only'):
        for i, row in enumerate(tables[-1].values):
            assert row == [i+1, X[i], result['Fitted X'][i], Y[i], result['Fitted Y'][i],
                           result['Residuals'][i], result['St Resid'][i]]
    elif name.startswith('fit_'):
        parameters = list(result['fits'])
        if name == 'fit_lognorm':
            parameters[0] = math.log(parameters[0])
        assert [row[1] for row in tables[0].values] == parameters
        assert tables[1].values == [list(result['ABL'])]
    elif name == 'compare':
        assert tables[-1].values == [[r[0], r[1][1], r[-1], r[1][2], r[1][0]] for r in result]
    for index, table in enumerate(tables):
        assert len(table.values) == len(table.display)
        xlsx = tmp_path / f'{index}.xlsx'
        png = tmp_path / f'{index}.png'
        export_excel(table, xlsx)
        export_png(table, png)
        book = load_workbook(xlsx)
        sheet = book.active
        assert sheet.max_row == len(table.values)+2
        assert sheet.max_column == len(table.headers)
        for r, row in enumerate(table.values, 3):
            for c, raw in enumerate(row, 1):
                cell = sheet.cell(r, c)
                if isinstance(raw, (int, float, np.number)) and math.isfinite(raw):
                    assert cell.data_type == 'n'
                    assert cell.value == pytest.approx(raw, rel=1e-14, abs=0)
                    assert cell.number_format == table.number_formats[c-1]
                else:
                    assert cell.value == table.display[r-3][c-1]
        book.close()
        with Image.open(png) as image:
            assert image.height == 72+36*(len(table.values)+1)
            assert image.crop((0, image.height-45, image.width, image.height)).getextrema() != ((255, 255),)*3
            image.load()


@pytest.mark.parametrize('model_name', ['OLS', 'GLS', 'Logit'])
def test_native_models_keep_text_and_unrounded_values(model_name, tmp_path, viewer, monkeypatch):
    import statsmodels.api as sm
    from pandastable_local.stats import StatsViewer, model_summary_tables
    from stat_table import export_excel, export_png
    x = np.arange(30.)
    y = np.array([0, 1, 0, 0, 1]*6) if model_name == 'Logit' else 1.23456789+x*1e-6+np.sin(x)*1e-8
    model = getattr(sm, model_name)(y, sm.add_constant(x))
    result = model.fit(disp=False) if model_name == 'Logit' else model.fit()
    summary = result.summary()
    monkeypatch.setattr(result, 'summary', lambda: summary)
    tables = list(model_summary_tables(result, summary))
    assert '\n'.join(map(str, tables)) + ('\n\n'+summary.extra_txt if summary.extra_txt else '') == str(summary)
    np.testing.assert_array_equal([row[1] for row in tables[1].values], result.params)
    np.testing.assert_array_equal([row[4] for row in tables[1].values], result.pvalues)
    np.testing.assert_array_equal([row[5:] for row in tables[1].values], result.conf_int())
    if model_name != 'Logit':
        assert any('e-' in cell for row in tables[1].display for cell in row)
    for i, table in enumerate(tables):
        export_excel(table, tmp_path / f'{i}.xlsx')
        export_png(table, tmp_path / f'{i}.png')
        with Image.open(tmp_path / f'{i}.png') as image:
            assert image.height == 72+36*(len(table.values)+1)
            image.verify()
    book = load_workbook(tmp_path / '1.xlsx')
    assert book.active['B3'].value == pytest.approx(result.params[0], rel=1e-14)
    assert book.active['B3'].data_type == 'n'
    book.close()
    StatsViewer.summary(SimpleNamespace(fit=result, fitinfo=viewer))
    viewer.main.update()
    assert viewer.T.get('1.0', 'end-1c') == str(summary)
    assert len(viewer.table_blocks) == len(tables)
    viewer.save(str(tmp_path / 'native.txt'))
    assert (tmp_path / 'native.txt').read_text(encoding='utf-8') == str(summary)


def test_real_app_routes_all_reports_to_g8_and_preserves_txt(viewer, tmp_path, monkeypatch):
    from MiniJMP import Minijmp
    from stat_table import StatTable
    import stat_table
    app = SimpleNamespace(table=SimpleNamespace(tOut=viewer), tOut=viewer)
    outputs = []
    for call in reports().values():
        call(outputs.append)
    for output in outputs:
        Minijmp.print(app, output)
    viewer.main.update()
    assert len(viewer.table_blocks) == sum(isinstance(o, StatTable) for o in outputs)
    expected = ''.join(str(o)+'\n' for o in outputs)
    assert viewer.T.get('1.0', 'end-1c') == expected
    viewer.save(str(tmp_path / '完整报告.txt'))
    assert (tmp_path / '完整报告.txt').read_text(encoding='utf-8') == expected
    block = viewer.table_blocks[-1]
    for menu_index, suffix in [(0, '.xlsx'), (1, '.png')]:
        filename = tmp_path / ('菜单导出'+suffix)
        monkeypatch.setattr(stat_table.filedialog, 'asksaveasfilename', lambda **k: str(filename))
        block.menu.invoke(menu_index)
        assert filename.stat().st_size > 100
    viewer.clear()
    assert not viewer.table_blocks and viewer.T.get('1.0', 'end-1c') == ''


def variable(value):
    return SimpleNamespace(get=lambda: value)


@pytest.mark.parametrize('multi', [False, True])
def test_correlation_dialogs_reach_all_four_methods(multi):
    import pandas as pd
    from stat_table import StatTable
    outputs = []
    figure = Figure()
    dialog = SimpleNamespace(df=pd.DataFrame({'温度': X, '压力': Y}),
        xvar=variable('温度'), yvar=variable('压力'), grpcols=['温度', '压力'],
        error=Mock(), get_plot_settings=lambda: None,
        app=SimpleNamespace(print=outputs.append, showPlotViewer=lambda **k: SimpleNamespace(fig=figure)))
    for key, value in dict(cor_p=True, cor_s=True, cor_k=True, cor_h=True, ellipse=True,
            red_ellipse=False, scatter=True, alpha=.05, ax_margin=.1, ylabel_0deg=False,
            xlabel_45deg=False, xtick_45deg=False).items():
        setattr(dialog, key, variable(value))
    for key, value in dict(xlabel=None, ylabel=None, x_min=None, x_max=None,
                           y_min=None, y_max=None, show_legend=False, grid=False).items():
        setattr(dialog, key, value)
    if not multi:
        dialog.ax_margin = .1
    (fit_dialog.MCorDialog if multi else fit_dialog.CorrelationDialog).apply(dialog)
    dialog.error.assert_not_called()
    tables = [o for o in outputs if isinstance(o, StatTable)]
    assert len(tables) == (8 if multi else 5)
    assert all(any(method in table.title for table in tables)
               for method in ['Pearson', 'Spearman', 'Kendall', 'Hoeffding'])


def test_g8_scientific_format_and_special_text_are_preserved(tmp_path):
    from stat_table import StatTable, export_excel
    table = StatTable(title='精度', headers=['变量', '极小 p 值'],
        values=[['=温度', 1.234567890123456e-100], ['缺失', None], ['标记', '< 1e-100']],
        display=[['=温度', '1.235e-100'], ['缺失', '*'], ['标记', '< 1e-100']],
        number_formats=['General', '0.000E+00'])
    export_excel(table, tmp_path / 'science.xlsx')
    book = load_workbook(tmp_path / 'science.xlsx')
    sheet = book.active
    assert sheet['A3'].data_type == 's' and sheet['A3'].value == '=温度'
    assert sheet['B3'].value == pytest.approx(1.234567890123456e-100, rel=1e-14, abs=0)
    assert sheet['B3'].number_format == '0.000E+00'
    assert sheet['B4'].value == '*' and sheet['B5'].value == '< 1e-100'
    book.close()


def test_long_residual_export_contains_last_observation(tmp_path):
    from stat_table import StatTable, export_excel, export_png
    x = np.arange(1., 101.)
    y = 2*x+np.sin(x)
    outputs = []
    result = ortho_fit.orthogonal_fit(x, y, print_residual=True, print_port=outputs.append)
    table = [o for o in outputs if isinstance(o, StatTable)][-1]
    export_excel(table, tmp_path / 'residuals.xlsx')
    book = load_workbook(tmp_path / 'residuals.xlsx')
    assert book.active.max_row == 102 and book.active.max_column == 7
    assert book.active['A102'].value == 100
    assert book.active['G102'].value == pytest.approx(result['St Resid'][-1], rel=1e-14)
    book.close()
    export_png(table, tmp_path / 'residuals.png')
    with Image.open(tmp_path / 'residuals.png') as image:
        assert image.height == 3708
        assert image.crop((0, 3660, image.width, 3708)).getextrema() != ((255, 255),)*3
        image.load()


def test_native_summary_window_can_close_and_reopen(tk_root):
    import statsmodels.api as sm
    from pandastable_local.stats import StatsViewer
    result = sm.OLS(Y, sm.add_constant(X)).fit()
    owner = SimpleNamespace(fit=result, fitinfo=None, parent=tk_root)
    for _ in range(2):
        StatsViewer.summary(owner)
        owner.w.update()
        assert len(owner.fitinfo.table_blocks) == 3
        close_command = owner.w.protocol('WM_DELETE_WINDOW')
        owner.w.tk.call(close_command)
        assert owner.fitinfo is None


@pytest.mark.parametrize('filename,names', [
    ('linear_fit.py', ['linear_fit']), ('ortho_fit.py', ['orthogonal_fit']),
    ('nonpara_cor.py', ['print_correlation']), ('dist_fit.py', ['print_table', 'compare_dist']),
    ('fit_dialog.py', ['print_multi_col_tables']),
    ('linear_plot.py', ['multi_fit', 'multi_fit_orig']),
    ('utilities.py', ['pearson_correlation', 'norm_test']),
])
def test_scoped_report_ports_do_not_stringify_tables(filename, names):
    source = (Path(__file__).parents[1] / 'code' / filename).read_text(encoding='utf-8')
    for function in ast.parse(source).body:
        if not isinstance(function, ast.FunctionDef) or function.name not in names:
            continue
        ports = [call for call in ast.walk(function) if isinstance(call, ast.Call)
                 and isinstance(call.func, ast.Name) and call.func.id in ('print', 'print_port')]
        assert any(isinstance(arg, ast.Call) and isinstance(arg.func, ast.Name)
                   and arg.func.id == 'StatTable' for call in ports for arg in call.args)
        assert not any(isinstance(arg, ast.Call) and isinstance(arg.func, ast.Name)
                       and arg.func.id == 'str' for call in ports for arg in call.args)


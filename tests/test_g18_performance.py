"""G18：调用数量、原始数据、真实 Artist 与报告/导出协议。"""
import importlib.util
from copy import copy
from pathlib import Path
from unittest.mock import Mock

import matplotlib
matplotlib.use('Agg')
from matplotlib.figure import Figure
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import pytest

import linear_plot
import control_chart
import stat_table
import time_series
import time_series_dialog
from test_g13_time_series import dialog, data, choose

BEFORE = Path(__file__).resolve().parents[1] / 'docs/g18-evidence/before'


def original(name):
    spec = importlib.util.spec_from_file_location('g18_before_'+name, BEFORE/(name+'.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize('entry', ['multi_fit', 'multi_fit_orig'])
@pytest.mark.parametrize('ellipses', [True, False])
def test_matrix_reuses_directional_pearson_and_only_builds_used_axes(monkeypatch, entry, ellipses):
    rng = np.random.default_rng(18)
    values = [rng.normal(size=40)*scale + shift for scale, shift in [(1, 3), (-2, 5), (8, -1), (3, 20)]]
    alpha = .17
    expected = [[linear_plot.pearson_correlation(x, y, alpha=alpha, print_out=False)
                 for x in values] for y in values]
    pearson = Mock(wraps=linear_plot.pearson_correlation)
    monkeypatch.setattr(linear_plot, 'pearson_correlation', pearson)
    tables = Mock(wraps=linear_plot.gen_table)
    monkeypatch.setattr(linear_plot, 'gen_table', tables)
    import matplotlib.pyplot as plt
    try:
        result = getattr(linear_plot, entry)(values, list('ABCD'), alpha=alpha,
                                            ellipse=ellipses, red_ellipse=ellipses)
        fig = plt.gcf()
        axes = fig.axes[0].child_axes if entry == 'multi_fit' else fig.axes
        assert len(axes) == 6
        assert pearson.call_count == 16
        assert tables.call_count == 0
        for y in range(4):
            for x in range(4):
                for key in ('r', 'p', 'r_l', 'r_u', 'cov', 'N'):
                    assert result['pearson'][y][x][key] == pytest.approx(expected[y][x][key])
        for ax in axes:
            points = ax.collections[0].get_offsets()
            x, y = next((x, y) for x in range(3) for y in range(x+1, 4)
                        if np.array_equal(points, np.column_stack([values[x], values[y]])))
            if ellipses:
                ellipse = expected[y][x]['ellipse']
                angle = np.linspace(0, 2*np.pi, 100)
                a, b, theta = (ellipse[k] for k in ('a', 'b', 'theta'))
                np.testing.assert_allclose(ax.lines[0].get_xdata(), np.mean(values[x])+
                    a*np.cos(angle)*np.cos(theta)-b*np.sin(angle)*np.sin(theta))
                np.testing.assert_allclose(ax.lines[0].get_ydata(), np.mean(values[y])+
                    a*np.cos(angle)*np.sin(theta)+b*np.sin(angle)*np.cos(theta))
                np.testing.assert_allclose(ax.patches[0].get_xy()[:100],
                    np.column_stack(ax.lines[0].get_data()))
        fig.canvas.draw()
    finally:
        plt.close('all')


def test_inset_grid_keeps_triangle_positions():
    fig = Figure()
    root = fig.subplots()
    axes = linear_plot.create_inset_grid(root, 3)
    for y, row in enumerate(axes):
        for x, ax in enumerate(row):
            if x > y:
                assert ax is None
            else:
                np.testing.assert_allclose(ax.get_axes_locator()(ax, None).bounds,
                    [root.get_position().x0+x*root.get_position().width/3,
                     root.get_position().y0+(2-y)*root.get_position().height/3,
                     root.get_position().width/3, root.get_position().height/3])


@pytest.mark.parametrize('entry', ['multi_fit', 'multi_fit_orig'])
def test_matrix_reports_keep_original_txt(entry):
    import matplotlib.pyplot as plt
    values = [[1., 2., 3., 4., 5.], [4., 8., 2., 9., 6.], [3., 2., 8., 4., 7.]]
    old, new = [], []
    try:
        getattr(original('linear_plot'), entry)(values, ['温度', '压力', '=流量'], print_out=True, print_port=old.append)
        getattr(linear_plot, entry)(values, ['温度', '压力', '=流量'], print_out=True, print_port=new.append)
        assert new == old
        for before, after in zip(old, new):
            if isinstance(after, stat_table.StatTable):
                assert before.values == after.values
                assert before.display == after.display
    finally:
        plt.close('all')


def test_spc_report_keeps_full_text_values_and_formats(monkeypatch):
    outputs = []
    constructed = []
    charts = [('I', [1., np.nan, 5., 8.], [0., 0., 0., 0.], 3., [7., 7., 9., 9.]),
              ('mR', [np.nan, 2., 3., 1.], np.array([np.nan]*4), 2., 5.)]
    for module in (original('control_chart'), control_chart):
        report = []
        tables = []
        gen_table = module.gen_table

        def capture(rows):
            table = gen_table(rows)
            tables.append(table)
            return table

        with monkeypatch.context() as patch:
            patch.setattr(module, 'gen_table', capture)
            module._print_spc_report('测试', charts, [1, 2, 2, 3], 'N', [101, '乙'], ['说明'], report.append)
        outputs.append(report)
        constructed.append(tables)
    assert outputs[0] == outputs[1]
    for before, after in zip(*constructed):
        assert before.rows == after.rows
        assert [[type(cell) for cell in row] for row in before.rows] == [
            [type(cell) for cell in row] for row in after.rows]
    for before, after in zip(*outputs):
        if isinstance(after, stat_table.StatTable):
            assert before.display == after.display
            np.testing.assert_equal(before.values, after.values)
            assert before.number_formats == after.number_formats


def test_spc_formats_scalar_limits_once(monkeypatch):
    isnan = Mock(wraps=np.isnan)
    monkeypatch.setattr(control_chart.np, 'isnan', isnan)
    control_chart._print_spc_report('I', [('I', list(range(100)), 0., 50., 100.)],
                                    1, 'N', None, [], lambda value: None)
    assert isnan.call_count <= 120


def test_excel_preserves_original_cells_styles_widths_and_precision(tmp_path):
    table = stat_table.StatTable(title='=标题', headers=['=标签', 'N', '数值', '比率'],
        values=[['=1+1', 123, 1.23456789012345, .123456789012345],
                ['@标签', 0, None, np.nan], ['末行', 2, np.inf, -np.inf]],
        number_formats=['General', '0', '0.00000000000000', '0.00%'])
    books = []
    try:
        for index, module in enumerate((original('stat_table'), stat_table)):
            path = tmp_path/f'{index}.xlsx'
            module.export_excel(table, path)
            books.append(load_workbook(path))
        before, after = [book.active for book in books]
        assert before.max_row == after.max_row == 5
        assert before.max_column == after.max_column == 4
        assert before.freeze_panes == after.freeze_panes == 'B3'
        for left_row, right_row in zip(before, after):
            for left, right in zip(left_row, right_row):
                for attribute in ('value', 'data_type', 'number_format', 'font', 'alignment'):
                    assert copy(getattr(left, attribute)) == copy(getattr(right, attribute)), (left.coordinate, attribute)
        for column in before.column_dimensions:
            assert before.column_dimensions[column].width == after.column_dimensions[column].width
        assert after['A1'].data_type == after['A2'].data_type == after['A3'].data_type == 's'
        assert after['C3'].value == 1.23456789012345
    finally:
        for book in books:
            book.close()


def test_excel_save_error_does_not_leave_streaming_temp_files(tmp_path):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    table = stat_table.StatTable(title='错误导出', headers=['变量', 'N'], values=[['温度', 123]])
    before = list(ALL_TEMP_FILES)
    with pytest.raises(OSError):
        stat_table.export_excel(table, tmp_path/'不存在的目录'/'report.xlsx')
    assert list(ALL_TEMP_FILES) == before


@pytest.mark.parametrize('dtype', ['float64', 'int64', 'Float64', 'Int64'])
def test_prepare_keeps_native_arrays_and_exact_nullable_values(dtype):
    values = [2**60+1, 2**60+3] if 'int' in dtype.lower() else [1.25, np.nan]
    data = pd.DataFrame({'Y': pd.Series(values, dtype=dtype)})
    _, ys = time_series.prepare_time_series(data, ['Y'])
    expected = data['Y'].to_numpy(dtype=object, na_value=np.nan)
    missing = pd.isna(expected)
    np.testing.assert_equal(pd.isna(ys['Y']), missing)
    assert list(ys['Y'][~missing]) == list(expected[~missing])
    if dtype in ('float64', 'int64'):
        assert ys['Y'].dtype == data['Y'].dtype


def test_dialog_prepares_once_and_revalidates_next_apply(dialog, monkeypatch):
    choose(dialog)
    prepare = Mock(wraps=time_series.prepare_time_series)
    monkeypatch.setattr(time_series, 'prepare_time_series', prepare)
    monkeypatch.setattr(time_series_dialog, 'prepare_time_series', prepare)
    dialog.apply()
    assert prepare.call_count == 1
    old_window = dialog.app.table.pf.main
    dialog.df.loc[:, '温度'] = np.nan
    dialog.apply()
    assert prepare.call_count == 2
    assert old_window.winfo_exists()
    assert any('全部缺失' in str(s) for s in dialog.output)


def test_vendor_viewer_creates_real_curves_all_current_paths(tk_root):
    import tkinter as tk
    from modified_table import MTable
    from pandastable_local.plotting import PlotViewer
    frame = tk.Frame(tk_root)
    frame.pack()
    data = pd.DataFrame({'A': [1., 3., np.nan, 7.], 'B': [8., 5., 4., 2.]})
    table = MTable(frame, dataframe=data, showtoolbar=False, showstatusbar=False)
    table.show()
    tk_root.update()
    viewer = PlotViewer(table=table, showoptions=True)
    try:
        for index, plot in enumerate((lambda: viewer.replot(data=data.copy()), viewer.plotCurrent, viewer.plot2D)):
            if index == 2:
                viewer._initFigure()  # 与 plotCurrent 的生产调用顺序一致。
            plot()
            viewer.main.update()
            assert len(viewer.ax.lines) == 2
            for line, column in zip(viewer.ax.lines, data):
                np.testing.assert_equal(line.get_ydata(), data[column].to_numpy())
            assert np.isnan(viewer.ax.lines[0].get_path().vertices[:, 1]).sum() == 1
        assert viewer.getcmap('viridis').name == 'viridis'
        assert viewer.getcmap('invalid cmap').name == 'Spectral'
    finally:
        viewer.close()
        table.destroy()
        frame.destroy()

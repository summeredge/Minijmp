"""REVIEW-008/013：以原始标量及真实 XLSX 单元格为来源参考。"""
from datetime import datetime
from hashlib import sha256

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from pca_analysis import fit_pca, load_pca
from test_worksheets import workbook
from test_csv_import import csv_app


@pytest.mark.parametrize('labels', [
    [9007199254740993, None, 9007199254740995, 9007199254740997],
    ['001', '002', '001', None],
])
def test_pca_source_labels_fit_project_tables_and_json(labels, tmp_path):
    df = pd.DataFrame({'x': [1., 3., 2., 8.], 'y': [5., 2., 9., 4.]},
                      index=pd.Index(labels, dtype=object))
    model = fit_pca(df, ['x', 'y'], 1)
    path = tmp_path / 'model.json'
    model.save(path)
    restored = load_pca(path)
    projected_df = df.copy()
    projected_df.iloc[1, 0] = np.nan
    expected_projected = [labels[i] for i in (0, 2, 3)]
    for candidate in (model, restored):
        for result, expected, rows in (
            (candidate.training, labels, [1, 2, 3, 4]),
            (candidate.project(projected_df, missing='drop'), expected_projected, [1, 3, 4]),
        ):
            for table in [result.source, *result.tables.values()]:
                actual = table['source_index'].tolist()
                assert actual == expected
                assert [type(v) for v in actual] == [type(v) for v in expected]
                assert table.source_row.tolist() == rows
                assert table.source_index.dtype == object


def test_excel_text_numbers_dates_missing_and_sheets(workbook, tmp_path):
    app, errors = workbook
    book = Workbook()
    first = book.active
    first.title = '标识'
    first.append(['ID', '数值', '日期', '混合', '缺失'])
    first.append(['001', 10, datetime(2026, 1, 2), '003', None])
    first.append(['002', 20.5, datetime(2026, 1, 3), 4, 8])
    second = book.create_sheet('第二张')
    second.append(['ID', '数值'])
    second.append(['0007', 7])
    path = tmp_path / 'text_ids.xlsx'
    book.save(path)
    original = sha256(path.read_bytes()).digest()
    assert app.load_excel(path)
    frame = app.sheets['标识'].model.df
    assert frame.ID.tolist() == ['001', '002']
    assert frame['数值'].tolist() == [10, 20.5]
    assert pd.api.types.is_numeric_dtype(frame['数值'])
    assert frame['日期'].tolist() == [pd.Timestamp(2026, 1, 2), pd.Timestamp(2026, 1, 3)]
    assert frame['混合'].tolist() == ['003', 4]
    assert pd.isna(frame['缺失'].iloc[0]) and frame['缺失'].iloc[1] == 8
    assert app.sheets['第二张'].model.df.ID.tolist() == ['0007']
    assert sha256(path.read_bytes()).digest() == original
    errors.assert_not_called()

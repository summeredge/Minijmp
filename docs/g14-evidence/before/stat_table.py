"""可复用统计表格；str(table) 保留 print_port / TXT 的文本协议。"""
import math
from numbers import Real
from tkinter import ttk, filedialog, messagebox
import tkinter as tk

RESULT_BACKGROUND = '#FFFFFF'


def configure_result_styles(style):
    """Keep result content white without changing application-wide styles."""
    for name in ('Result.TFrame', 'Result.TLabel', 'Result.TMenubutton'):
        style.configure(name, background=RESULT_BACKGROUND)
        style.map(name, background=[('', RESULT_BACKGROUND)])


def display_value(value, column):
    if isinstance(value, Real):
        if not math.isfinite(value):
            return '*'
        return str(int(value)) if column in ('N', 'N*', '众数频数') else f'{value:.3f}'
    return '*' if value is None else str(value)


class StatTable(str):
    """标题、列名、二维原值及显示字符串的共同载体。

    兼容 StatTable(text, title, headers, values)。省略 text 时由显示数据生成
    PrettyTable；传入 text（含 PrettyTable 对象）时原样保留其字符串。
    number_formats 为逐列 Excel 格式；不改变原始数值精度或 display。
    """

    def __new__(cls, text=None, title='', headers=(), values=(), *,
                display=None, number_formats=None):
        if isinstance(headers, (str, bytes)):
            raise TypeError('列名必须为字符串序列，不能为单个字符串')
        headers = list(headers)
        if not headers or not all(isinstance(header, str) for header in headers):
            raise ValueError('列名必须为非空的字符串序列')
        if not isinstance(title, str):
            raise TypeError('标题必须为字符串')
        values = list(values)
        if any(isinstance(row, (str, bytes)) for row in values):
            raise ValueError('行数据必须为二维序列')
        values = [list(row) for row in values]
        if any(len(row) != len(headers) for row in values):
            raise ValueError('每行数据的列数必须与列名一致')
        if display is None:
            display = [[display_value(v, h) for h, v in zip(headers, row)]
                       for row in values]
        else:
            display = list(display)
            if any(isinstance(row, (str, bytes)) for row in display):
                raise ValueError('显示数据必须为二维序列')
            display = [list(row) for row in display]
            if len(display) != len(values) or any(len(row) != len(headers) for row in display):
                raise ValueError('显示数据的行列数必须与原值一致')
            if any(not isinstance(cell, str) for row in display for cell in row):
                raise TypeError('显示数据的单元格必须为字符串')
        if number_formats is None:
            number_formats = ['0' if h in ('N', 'N*', '众数频数') else '0.000' for h in headers]
        else:
            if isinstance(number_formats, (str, bytes)):
                raise TypeError('数字格式必须为逐列序列')
            number_formats = list(number_formats)
            if len(number_formats) != len(headers) or any(
                    not isinstance(fmt, str) or not fmt for fmt in number_formats):
                raise ValueError('数字格式必须为与列名等长的非空字符串序列')
        if text is None:
            from prettytable import PrettyTable
            text = PrettyTable(headers)
            text.add_rows(display)
            text.align = 'r'
            text.align[headers[0]] = 'l'
        obj = super().__new__(cls, str(text))
        obj.title = title
        obj.headers = headers
        obj.values = values
        obj.display = display
        obj.number_formats = number_formats
        return obj


def export_excel(table, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = '统计结果'
    sheet.append([table.title])
    sheet.append(table.headers)
    for row, shown in zip(table.values, table.display):
        sheet.append([v if isinstance(v, Real) and math.isfinite(v) else text
                      for v, text in zip(row, shown)])
    for row in sheet:
        for cell in row:
            cell.font = Font(name='Microsoft YaHei', bold=cell.row <= 2)
            cell.alignment = Alignment(horizontal='right' if cell.column > 1 else 'left')
            if cell.data_type == 'f':  # 标签永远作为文本，避免 Excel 将其解释为公式。
                cell.data_type = 's'
            if cell.row > 2 and cell.data_type == 'n':
                cell.number_format = table.number_formats[cell.column-1]
    for i, header in enumerate(table.headers, 1):
        sheet.column_dimensions[get_column_letter(i)].width = max(12, len(header)*2+2,
            max((len(row[i-1])*2+2 for row in table.display), default=0))
    sheet.freeze_panes = 'B3'
    book.save(filename)
    book.close()


def export_png(table, filename):
    from PIL import Image, ImageDraw, ImageFont
    from matplotlib import font_manager
    path = font_manager.findfont(font_manager.FontProperties(
        family=['Microsoft YaHei', 'SimHei', 'Noto Sans CJK SC', 'DejaVu Sans']))
    face = ImageFont.truetype(path, 20)
    rows = [table.headers] + table.display
    widths = [max(face.getlength(row[i]) for row in rows) + 32
              for i in range(len(table.headers))]
    width = int(max(sum(widths)+32, face.getlength(table.title)+40))
    image = Image.new('RGB', (width, 72+36*len(rows)), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((16, 12), table.title, font=face, fill='#17476b')
    for r, row in enumerate(rows):
        x = 16
        for c, text in enumerate(row):
            offset = widths[c]-16-face.getlength(text) if c and r else 8
            draw.text((x+offset, 60+36*r), text, font=face, fill='#202020')
            x += widths[c]
    draw.line((16, 90, width-16, 90), fill='#9caab5', width=1)
    image.save(filename, format='PNG')


class TableBlock(ttk.Frame):
    def __init__(self, parent, table):
        configure_result_styles(ttk.Style(parent))
        super().__init__(parent, padding=8, style='Result.TFrame')
        self.table = table
        titlebar = ttk.Frame(self, style='Result.TFrame')
        titlebar.grid(row=0, column=0, sticky='w')
        self.menu_button = button = ttk.Menubutton(titlebar, text='导出', style='Result.TMenubutton')
        button.pack(side='left')
        self.title_label = ttk.Label(titlebar, text=table.title, style='Result.TLabel',
                                    font=('Microsoft YaHei', 11, 'bold'))
        self.title_label.pack(side='left', padx=(8, 0))
        self.menu = tk.Menu(button, tearoff=False)
        for kind in ('Excel', 'PNG'):
            self.menu.add_command(label='导出 '+kind, command=lambda k=kind: self.export(k))
        button.configure(menu=self.menu)
        body = ttk.Frame(self, style='Result.TFrame')
        body.grid(row=1, column=0, sticky='w')
        for r, row in enumerate([table.headers] + table.display):
            for c, text in enumerate(row):
                ttk.Label(body, text=text, anchor='e' if c else 'w', style='Result.TLabel',
                          font=('Microsoft YaHei', 10, 'bold' if r == 0 else 'normal')).grid(
                    row=r, column=c, sticky='ew', padx=10, pady=3)

    def export(self, kind):
        extension = '.xlsx' if kind == 'Excel' else '.png'
        filename = filedialog.asksaveasfilename(parent=self, defaultextension=extension,
            filetypes=[(kind, '*'+extension)])
        if not filename:
            return
        try:
            (export_excel if kind == 'Excel' else export_png)(self.table, filename)
        except Exception as error:
            messagebox.showerror('导出失败', str(error), parent=self)

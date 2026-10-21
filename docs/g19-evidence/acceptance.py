"""G19 重放 G18 真实 Tk/导出验收；仅改证据输出与历史源码定位。"""
from pathlib import Path
import sys
import tkinter as tk
from types import SimpleNamespace
import importlib.util
import json

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT/'code'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import ImageGrab
from openpyxl import load_workbook
from MiniJMP import Minijmp
from modified_app import Minijmp_pre
from modified_table import MTable
from stat_table import StatTable, export_excel, export_png
from txt_output import txt_viewer
from time_series_dialog import TimeSeriesDialog
from time_series import AXIS_MODES
from pandastable_local.plotting import PlotViewer
import linear_plot


def capture(window, name, geometry):
    window.geometry(geometry)
    window.lift()
    window.update()
    window.after(150)
    with ImageGrab.grab(window=window.winfo_id()) as image:
        image.save(OUT/name)


def main():
    evidence = {}
    spec = importlib.util.spec_from_file_location('before_linear_plot', ROOT/'docs/g18-evidence/before/linear_plot.py')
    before = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(before)
    rng = np.random.default_rng(18)
    x = rng.normal(size=100)
    values = [x, -2*x+rng.normal(size=100), rng.normal(size=100)*5, .5*x+rng.normal(size=100)]
    for name in ('multi_fit', 'multi_fit_orig'):
        for phase, module in [('before', before), ('after', linear_plot)]:
            getattr(module, name)(values, ['温度 ℃', '压力 MPa', '流量 m³/h', '浓度 %'], alpha=.1)
            fig = plt.gcf()
            fig.set_size_inches(9, 8)
            fig.canvas.draw()
            fig.savefig(OUT/f'{name}-{phase}.png', dpi=100)
            evidence[f'{name}-{phase}-axes'] = len(fig.axes)+sum(len(a.child_axes) for a in fig.axes)
            plt.close(fig)

    root = tk.Tk()
    root.withdraw()
    callback_errors = []
    root.report_callback_exception = lambda exc, value, tb: callback_errors.append(str(value))
    viewer = txt_viewer(SimpleNamespace(tOut=None))
    app = SimpleNamespace(table=viewer.table, tOut=viewer)
    values = [[f'=样本 {i+1}', i+1, i+1/7, None if i == 3 else i/3000] for i in range(1003)]
    table = StatTable(title='完整导出与末页：1,003 行', headers=['=标签', 'N', '测量值', '比例'],
        values=values, number_formats=['General', '0', '0.00000000000000', '0.00%'])
    Minijmp.print(app, '报告开始')
    Minijmp.print(app, table)
    Minijmp.print(app, '报告结束')
    expected = '报告开始\n'+str(table)+'\n报告结束\n'
    viewer.save(OUT/'full-report.txt')
    assert (OUT/'full-report.txt').read_text(encoding='utf-8') == expected
    copied = []
    viewer.T.clipboard_clear = lambda: None
    viewer.T.clipboard_append = copied.append
    viewer.select_all()
    viewer.copy()
    assert copied == [expected]
    viewer.T.tag_remove('sel', '1.0', 'end')
    block = viewer.table_blocks[0]
    capture(viewer.main, 'report-first-page.png', '800x650+60+60')
    block.show_page(block.page_count)
    assert block.tree.item(block.tree.get_children()[-1], 'values') == tuple(table.display[-1])
    capture(viewer.main, 'report-last-page.png', '800x650+60+60')
    capture(viewer.main, 'report-small-window.png', '480x320+60+60')
    export_excel(table, OUT/'full-report.xlsx')
    book = load_workbook(OUT/'full-report.xlsx')
    sheet = book.active
    assert (sheet.max_row, sheet.max_column, sheet.freeze_panes) == (1005, 4, 'B3')
    for raw, cells in zip(values, sheet.iter_rows(min_row=3)):
        assert cells[0].value == raw[0] and cells[0].data_type == 's'
        assert cells[1].value == raw[1]
        assert abs(cells[2].value-raw[2]) < 1e-11
        assert cells[2].number_format == '0.00000000000000'
        if raw[3] is None:
            assert cells[3].value == '*'
        else:
            assert abs(cells[3].value-raw[3]) < 1e-15
    book.close()
    try:
        export_png(table, OUT/'too-large.png')
    except ValueError as exc:
        evidence['png_limit'] = str(exc)
    else:
        raise AssertionError('大表 PNG 未受限')
    viewer.clear()
    small = StatTable(title='小表外观与 PNG', headers=['变量', '均值', 'N'],
                      values=[['温度', 1.23456789012345, 100], ['压力', 8.7654321, 100]])
    Minijmp.print(app, small)
    capture(viewer.main, 'small-report.png', '720x500+60+60')
    export_png(small, OUT/'small-report-export.png')
    viewer.to_end(defer=True)
    assert viewer._scroll_job is not None
    viewer.close()
    assert viewer._scroll_job is None and block.table is None
    evidence['report'] = {'rows': 1003, 'txt_exact': True, 'copy_exact': True,
                           'excel_all_rows_checked': True, 'close_cleanup': True}

    parent = tk.Frame(root)
    parent.pf = None
    app = SimpleNamespace(table=parent, print=print, addPlot=lambda: None,
        dialog_help_msg=lambda **kw: Minijmp_pre.dialog_help_msg(None, **kw))
    app.showPlotViewer = lambda **kw: Minijmp.showPlotViewer(app, **kw)
    t = pd.date_range('2026-10-21', periods=1000, freq='s').to_numpy().copy()
    t[100] = t[99]
    v = np.sin(np.arange(1000)/50)
    v[350:380] = np.nan
    data = pd.DataFrame({'采样时间': t[::-1], '温度': v, '压力': 2+np.cos(np.arange(1000)/90),
                         '流量': 500+np.arange(1000)})
    dlg = TimeSeriesDialog(parent, df=data, app=app, title='G18 时间序列')
    dlg.ylist.selection_set(0, 2)
    dlg.mode.set(AXIS_MODES[2])
    dlg.update_variables()
    dlg.assignments = dict(zip(dlg.selected_columns(), [1, 2, 3]))
    dlg.titlevar.set('完整原始趋势：重复时间与 NaN 断线')
    dlg.apply()
    plot = parent.pf
    order = np.argsort(t[::-1], kind='stable')
    assert len(dlg.selected_columns()) == len(plot.fig.axes) == 3
    assert all(len(ax.lines) == 1 for ax in plot.fig.axes)
    for ax, column in zip(plot.fig.axes, dlg.selected_columns()):
        np.testing.assert_equal(ax.lines[0].get_ydata(), data[column].to_numpy()[order])
        assert len(ax.lines[0].get_xdata()) == 1000
    capture(plot.main, 'time-series-window.png', '1000x750+60+60')
    plot.savePlot(str(OUT/'time-series.png'))
    evidence['time_series'] = {'points_per_curve': 1000, 'axes': 3, 'stable_sort_and_nan': True}
    plot.close()
    dlg.main.destroy()
    parent.destroy()

    frame = tk.Frame(root)
    vendor_data = data[['温度', '压力', '流量']]
    table = MTable(frame, dataframe=vendor_data, showtoolbar=False, showstatusbar=False)
    table.show()
    vendor = PlotViewer(table=table, showoptions=True)
    vendor.replot(data=vendor_data.copy())
    assert len(vendor.ax.lines) == 3
    for line, column in zip(vendor.ax.lines, vendor_data):
        np.testing.assert_equal(line.get_ydata(), vendor_data[column].to_numpy())
    capture(vendor.main, 'vendor-real-curves.png', '1000x700+60+60')
    vendor.savePlot(str(OUT/'vendor-real-curves-export.png'))
    vendor.close()
    table.destroy()
    frame.destroy()
    root.destroy()
    assert not callback_errors, callback_errors
    evidence['vendor_curves'] = 3
    evidence['tk_callback_errors'] = callback_errors
    (OUT/'gui-export-acceptance.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
